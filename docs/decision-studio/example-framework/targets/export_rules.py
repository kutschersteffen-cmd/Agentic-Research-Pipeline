"""Writes Logic_of_targets_decision_rules.xlsx and pseudocode.md next to framework.json.

Run from backend/: PYTHONPATH=. python3 ../docs/decision-studio/example-framework/targets/export_rules.py
"""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

T = Path(__file__).parent
cfg = json.loads((T / "framework.json").read_text())["config"]
F = "Arial"
HEAD = PatternFill("solid", fgColor="1F3864")
OUT = PatternFill("solid", fgColor="D9E1F2")
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
wb = Workbook()
wb.remove(wb.active)


def sheet(title, header, rows, widths, intro):
    ws = wb.create_sheet(title)
    for i, line in enumerate(intro, 1):
        ws.cell(i, 1, line).font = Font(name=F, bold=i == 1, size=12 if i == 1 else 10)
    r = len(intro) + 2
    for c, h in enumerate(header, 1):
        cell = ws.cell(r, c, h)
        cell.font, cell.fill, cell.border = Font(name=F, bold=True, color="FFFFFF"), HEAD, BOX
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = ws.cell(r + 1, 1)
    for row in rows:
        r += 1
        for c, v in enumerate(row, 1):
            cell = ws.cell(r, c, v)
            cell.font, cell.border = Font(name=F, size=10), BOX
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    return ws, r - len(rows) + 1


def dash(v):
    return v if v.strip() else "–"


def tables(graph):
    return [n for n in graph["nodes"] if n["type"] == "decisionTableNode"]


def exprs(graph):
    return [n for n in graph["nodes"] if n["type"] == "expressionNode"]


def rule_sentence(ins, outs, row):
    cond = " and ".join(f"{i['name']} = {row[i['id']]}" for i in ins if row[i["id"]].strip()) or "otherwise"
    return f"IF {cond} THEN " + ", ".join(f"{o['field']} = {row[o['id']]}" for o in outs)


for e in exprs(cfg["rule_graph"]):
    sheet(
        e["name"].capitalize(),
        ["Output column", "Expression"],
        [(x["key"], x["value"]) for x in e["content"]["expressions"]],
        [42, 140],
        [f"Rules tab > {e['name']} (expression node)", "Evaluated per company, top to bottom."],
    )

pseudo = ["# Logic of targets (Q2.1): pseudo code", "",
          "Generated from `framework.json`; same logic Decision Studio runs. Hit policy `first`: the first matching row wins; an empty input matches anything.", ""]
for tab, graph in (("Rules", cfg["rule_graph"]), ("Decision tree", cfg["tier_graph"])):
    for n in graph["nodes"]:
        if n["type"] == "expressionNode":
            pseudo += [f"## {tab} > {n['name']} (expressions)", "```"]
            pseudo += [f"{x['key']} = {x['value']}" for x in n["content"]["expressions"]] + ["```", ""]
        elif n["type"] == "decisionTableNode":
            c = n["content"]
            ins, outs = c["inputs"], c["outputs"]
            rules = [rule_sentence(ins, outs, r) for r in c["rules"]]
            pseudo += [f"## {tab} > {n['name']} (decision table, hit policy {c['hitPolicy']})", "```"]
            for k, r in enumerate(c["rules"]):
                cond = " and ".join(
                    f"{i['name']} has info" if "," in r[i["id"]] else f"{i['name']} {'' if r[i['id']][0] in '<>=!' else '= '}{r[i['id']]}"
                    for i in ins if r[i["id"]].strip()
                )
                then = ", ".join(f"{o['field']} = {r[o['id']]}" for o in outs)
                pseudo.append((("IF " if k == 0 else "ELSE IF ") + cond + " THEN " + then) if cond else "ELSE " + then)
            pseudo += ["```", ""]
            header = ["Rule"] + [f"IN: {i['name']}\n({i['field']})" for i in ins] + [f"OUT: {o['name']}\n({o['field']})" for o in outs] + ["Meaning"]
            rows = [[k] + [dash(r[i["id"]]) for i in ins] + [dash(r[o["id"]]) for o in outs] + [rule_sentence(ins, outs, r)] for k, r in enumerate(c["rules"], 1)]
            ws, first = sheet(n["name"], header, rows, [6] + [18] * len(ins) + [24] * len(outs) + [70],
                              [f"{tab} tab > {n['name']} (decision table, hit policy {c['hitPolicy']})", "First matching row wins; – matches anything."])
            for col in range(2 + len(ins), 2 + len(ins) + len(outs)):
                for r in range(first, first + len(rows)):
                    ws.cell(r, col).fill = OUT

wb.save(T / "Logic_of_targets_decision_rules.xlsx")
(T / "pseudocode.md").write_text("\n".join(pseudo) + "\n")
print(wb.sheetnames)
