"""Writes answers into a copy of the investor's own file (spec §5). Every answer is checked before anything changes,
so a refusal never leaves a half-filled file."""

from __future__ import annotations

import io
import posixpath
import re
import zipfile
from xml.etree import ElementTree

from arp.rfp.locations import docx_cell, paragraph_index, split_cell
from arp.rfp.text import is_placeholder

_NS = {
    "m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
}


class FillRefused(ValueError):
    """The file cannot be filled as asked; the message is one line naming the sheet or location."""


def fill(data: bytes, suffix: str, answers: list[tuple[str, str]]) -> bytes:
    """A copy of `data` with each (answer_location, answer) written in."""
    suffix = suffix.lower()
    if suffix == ".xlsx":
        return _fill_xlsx(data, answers)
    if suffix == ".docx":
        return _fill_docx(data, answers)
    raise FillRefused(f"Unsupported file type {suffix!r}: use .xlsx or .docx")


def _rels(z: zipfile.ZipFile, part: str) -> dict[str, tuple[str, str]]:
    """Relationship id -> (type, resolved part name) for `part`."""
    folder, name = posixpath.split(part)
    try:
        root = ElementTree.fromstring(z.read(posixpath.join(folder, "_rels", name + ".rels")))
    except KeyError:
        return {}
    out = {}
    for rel in root.findall("rel:Relationship", _NS):
        target = rel.get("Target", "")
        resolved = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join(folder, target))
        out[rel.get("Id")] = (rel.get("Type", ""), resolved)
    return out


def _sheet_with_drawing(data: bytes) -> str | None:
    """The first sheet (workbook order) that holds an image, chart or shape. openpyxl would drop it on save."""
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        if not any(re.fullmatch(r"xl/drawings/drawing[^/]*\.xml", n) for n in z.namelist()):
            return None
        sheets = ElementTree.fromstring(z.read("xl/workbook.xml")).findall("m:sheets/m:sheet", _NS)
        targets = _rels(z, "xl/workbook.xml")
        for sheet in sheets:
            part = targets.get(sheet.get(f"{{{_NS['r']}}}id"), ("", ""))[1]
            if any(t.endswith("/drawing") for t, _ in _rels(z, part).values()):
                return sheet.get("name")
        return sheets[0].get("name") if sheets else "?"


def _fill_xlsx(data: bytes, answers: list[tuple[str, str]]) -> bytes:
    from openpyxl import load_workbook
    from openpyxl.cell.cell import MergedCell

    sheet = _sheet_with_drawing(data)
    if sheet is not None:
        raise FillRefused(f"Sheet {sheet} holds an image, chart or shape that would be lost; answer this file by hand.")
    wb = load_workbook(io.BytesIO(data))
    cells = []
    for location, answer in answers:
        parts = split_cell(location)
        if parts is None or parts[0] not in wb.sheetnames:
            raise FillRefused(f"Answer location {location} does not exist.")
        cell = wb[parts[0]][parts[1]]
        if isinstance(cell, MergedCell) or not is_placeholder("" if cell.value is None else str(cell.value)):
            raise FillRefused(f"Answer location {location} is not empty or a placeholder.")
        cells.append((cell, answer))
    for cell, answer in cells:
        cell.value = answer
        cell.data_type = "s"  # an answer starting with "=" stays text, never a formula
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _set_text(paragraph, text: str) -> None:
    """The paragraph's text becomes `text`, in one run; its paragraph properties (style) stay."""
    from docx.oxml.ns import qn

    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    paragraph.add_run(text)


def _fill_docx(data: bytes, answers: list[tuple[str, str]]) -> bytes:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph

    doc = Document(io.BytesIO(data))
    paras = list(doc.paragraphs)  # by original index, so an inserted paragraph never shifts a later P<n>
    plan = []
    for location, answer in answers:
        n = paragraph_index(location)
        if n is not None:
            if n >= len(paras):
                raise FillRefused(f"Answer location {location} does not exist.")
            plan.append((n, None, answer))
            continue
        cell = docx_cell(doc, location)
        if cell is None:
            raise FillRefused(f"Answer location {location} does not exist.")
        if not is_placeholder(cell.text):
            raise FillRefused(f"Answer location {location} is not empty or a placeholder.")
        plan.append((None, cell, answer))
    for n, cell, answer in plan:
        if cell is not None:
            first = cell.paragraphs[0]
            for child in list(cell._tc):
                if child.tag != qn("w:tcPr") and child is not first._p:
                    cell._tc.remove(child)
            _set_text(first, answer)
            continue
        after = paras[n + 1] if n + 1 < len(paras) else None
        if after is not None and is_placeholder(after.text):
            _set_text(after, answer)
            continue
        new = OxmlElement("w:p")
        paras[n]._p.addnext(new)
        para = Paragraph(new, paras[n]._parent)
        if after is not None:
            para.style = after.style
        else:
            try:
                para.style = doc.styles["Normal"]
            except KeyError:
                pass  # no style called Normal: the document's default paragraph style applies
        para.add_run(answer)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
