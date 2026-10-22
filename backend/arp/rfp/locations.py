"""Locations in an RFP file: `<sheet>!<cell>` in Excel, `P<n>` and `T<t>R<r>C<c>` in Word (spec, Shared definitions)."""

from __future__ import annotations

import io
import re
import zipfile
from datetime import date
from xml.etree import ElementTree

SUFFIXES = (".xlsx", ".docx")
_CELL = re.compile(r"^([A-Z]{1,3})([1-9]\d{0,6})$")
_PARA = re.compile(r"^P(\d+)$")
_TABLE_CELL = re.compile(r"^T(\d+)R(\d+)C(\d+)$")
_MAX_ROW, _MAX_COL = 1_048_576, 16_384


def split_cell(location: str) -> tuple[str, str] | None:
    """(sheet, coordinate) for a well-formed Excel location, else None. The sheet is not checked."""
    from openpyxl.utils import column_index_from_string

    sheet, _, coord = location.rpartition("!")
    m = _CELL.match(coord)
    if not sheet or not m or int(m.group(2)) > _MAX_ROW or column_index_from_string(m.group(1)) > _MAX_COL:
        return None
    return sheet, coord


def paragraph_index(location: str) -> int | None:
    m = _PARA.match(location)
    return int(m.group(1)) if m else None


def table_cell_index(location: str) -> tuple[int, int, int] | None:
    m = _TABLE_CELL.match(location)
    return (int(m.group(1)), int(m.group(2)), int(m.group(3))) if m else None


def docx_cell(doc, location: str):
    """The top-level table cell at `T<t>R<r>C<c>`, or None when it does not exist."""
    idx = table_cell_index(location)
    if idx is None:
        return None
    t, r, c = idx
    if t >= len(doc.tables) or r >= len(rows := doc.tables[t].rows) or c >= len(cells := rows[r].cells):
        return None
    return cells[c]


class Layout:
    """Every location of a file with its text, in file order. `text(location)` is None when the location does not
    exist; an Excel cell that exists but is empty reads as ""."""

    def __init__(self, data: bytes, suffix: str) -> None:
        self.suffix = suffix.lower()
        if self.suffix == ".xlsx":
            from openpyxl import load_workbook

            wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
            self.sheets = list(wb.sheetnames)
            self.items = [
                (f"{name}!{cell.coordinate}", str(cell.value))
                for name in wb.sheetnames
                for row in wb[name].iter_rows()
                for cell in row
                if getattr(cell, "value", None) is not None and str(cell.value).strip()
            ]
        elif self.suffix == ".docx":
            from docx import Document
            from docx.oxml.ns import qn

            doc = Document(io.BytesIO(data))
            self.items = []
            p = t = 0
            for child in doc.element.body.iterchildren():
                if child.tag == qn("w:p"):
                    self.items.append((f"P{p}", doc.paragraphs[p].text))
                    p += 1
                elif child.tag == qn("w:tbl"):
                    for r, row in enumerate(doc.tables[t].rows):
                        self.items += [(f"T{t}R{r}C{c}", cell.text) for c, cell in enumerate(row.cells)]
                    t += 1
        else:
            raise ValueError(f"Unsupported file type {suffix!r}: use .xlsx or .docx")
        self._text = dict(self.items)
        self.order = {loc: i for i, (loc, _) in enumerate(self.items)}

    def text(self, location: str) -> str | None:
        if self.suffix == ".xlsx":
            cell = split_cell(location)
            if cell is None or cell[0] not in self.sheets:
                return None
            return self._text.get(location, "")
        return self._text.get(location)


def document_date(data: bytes) -> date | None:
    """The file's core property `modified`, else `created`, else None."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            root = ElementTree.fromstring(z.read("docProps/core.xml"))
    except (KeyError, zipfile.BadZipFile, ElementTree.ParseError):
        return None
    for name in ("modified", "created"):
        node = root.find(f"{{http://purl.org/dc/terms/}}{name}")
        if node is not None and node.text:
            try:
                return date.fromisoformat(node.text.strip()[:10])
            except ValueError:
                continue
    return None
