"""Golden checks with the real model (spec Constraints, AC-42..44). Case files describe fictitious RFPs as JSON and
are built into .xlsx/.docx in memory, so no binary or client file is committed."""

from __future__ import annotations

import io
import json
from datetime import date
from pathlib import Path

from pydantic import BaseModel, Field

from arp.llm.base import LLMClient
from arp.rfp.detect import detect_questions
from arp.rfp.draft import DRAFTED, NO_SOURCE, Entry, draft_answer
from arp.rfp.locations import Layout


def build_file(case: dict) -> tuple[bytes, str]:
    """(file bytes, suffix) for a case: `sheets` {name: rows} for Excel, `blocks` [{p, style} | {table: rows}] for Word."""
    out = io.BytesIO()
    if "sheets" in case:
        from openpyxl import Workbook

        wb = Workbook()
        wb.remove(wb.active)
        for name, rows in case["sheets"].items():
            ws = wb.create_sheet(name)
            for row in rows:
                ws.append(row)
        wb.save(out)
        return out.getvalue(), ".xlsx"
    from docx import Document

    doc = Document()
    for block in case["blocks"]:
        if "table" in block:
            rows = block["table"]
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            for r, row in enumerate(rows):
                for c, text in enumerate(row):
                    table.cell(r, c).text = text
        else:
            doc.add_paragraph(block["p"], style=block.get("style"))
    doc.save(out)
    return out.getvalue(), ".docx"


async def score_questions(cases: list[dict], llm: LLMClient) -> tuple[int, int, list[str]]:
    """(labelled questions kept with their labelled location and answer location, labelled total, misses)."""
    hit, total, misses = 0, 0, []
    for case in cases:
        data, suffix = build_file(case)
        found, _ = await detect_questions(Layout(data, suffix), llm)
        got = {(d.location, d.answer_location) for d in found}
        for label in case["labels"]:
            total += 1
            if (label["location"], label["answer_location"]) in got:
                hit += 1
            else:
                misses.append(f"{case['name']}: {label['location']} -> {label['answer_location']}")
    return hit, total, misses


class FactCheck(BaseModel):
    unsupported: list[str] = Field(default_factory=list, description="Each fact in the draft that no cited entry states.")


_VERIFY = """You check a drafted RFP answer against the library entries it cites. List every fact, figure, date or name
in the draft that the cited entries' answers do not state. Rewording, shortening and combining are fine. Return an empty
list when everything in the draft is stated in the entries."""


async def score_drafts(cases: dict, llm: LLMClient, verifier: LLMClient, *, today: date, settings=None) -> dict:
    """Counts for answerable questions drafted, unanswerable ones left `no source`, and drafts with unsupported facts."""
    entries = [Entry(i + 1, e["question"], e["answer"], date.fromisoformat(e.get("review_by", "2099-12-31")))
               for i, e in enumerate(cases["library"])]
    by_id = {e.id: e for e in entries}
    counts = dict(drafted=0, answerable=0, blank=0, unanswerable=0, unsupported=0)
    notes = []
    for q in cases["questions"]:
        d = await draft_answer(q["question"], entries, llm, today=today, settings=settings)
        if q["answerable"]:
            counts["answerable"] += 1
            counts["drafted"] += d.status == DRAFTED
        else:
            counts["unanswerable"] += 1
            counts["blank"] += d.status == NO_SOURCE
        if d.status != (DRAFTED if q["answerable"] else NO_SOURCE):
            notes.append(f"{d.status}: {q['question']}")
        if d.status == DRAFTED:
            cited = "\n\n".join(f"[E{i}] {by_id[i].answer}" for i in dict.fromkeys(s["entry_id"] for s in d.sources))
            check, _ = await verifier.complete_structured(
                system=_VERIFY, prompt=f"Cited entries:\n\n{cited}\n\nDraft: {d.answer}", output_model=FactCheck
            )
            if check.unsupported:
                counts["unsupported"] += 1
                notes.append(f"unsupported in '{d.answer}': {'; '.join(check.unsupported)}")
    return counts | {"notes": notes}


def load(path: Path):
    return json.loads(Path(path).read_text(encoding="utf-8"))
