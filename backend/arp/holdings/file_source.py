"""CSV and Excel holdings files, read through a JSON provider mapping. Files are untrusted: no formulas are evaluated."""

import csv
import hashlib
import io
import json
import re
from pathlib import Path
from typing import Literal

import openpyxl
from pydantic import BaseModel

from arp.decision.parsing import sniff_delimiter
from arp.holdings.validate import TEMPLATE_COLUMNS

MAPPINGS_DIR = Path(__file__).parent / "mappings"


class Mapping(BaseModel):
    provider: str
    columns: dict[str, str]
    decimal: Literal[".", ","] = "."
    weight_unit: Literal["percent", "fraction"] = "percent"
    sheet: str | None = None


def load_mapping(provider: str) -> Mapping:
    path = MAPPINGS_DIR / f"{provider}.json"
    if not re.fullmatch(r"[A-Za-z0-9_-]+", provider) or not path.is_file():
        raise ValueError(f"unknown mapping provider: {provider!r}")
    return Mapping.model_validate(json.loads(path.read_text(encoding="utf-8")))


def _cell(v):
    return v.strip() if isinstance(v, str) else v


MAX_ROWS = 200_000


def _records(data: bytes, suffix: str, mapping: Mapping, max_rows: int):
    if suffix == ".csv":
        text = data.decode("utf-8-sig")
        reader = csv.DictReader(io.StringIO(text, newline=""), delimiter=sniff_delimiter(text))
        yield from ((i, {k: v for k, v in rec.items() if k is not None}) for i, rec in enumerate(reader, start=2))
        return
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        ws = wb[mapping.sheet] if mapping.sheet else wb.worksheets[0]
        it = ws.iter_rows(values_only=True, max_row=max_rows + 2)
        header = [None if h is None else str(h).strip() for h in next(it, ())]
        yield from ((i, dict(zip(header, vals, strict=False))) for i, vals in enumerate(it, start=2))
    finally:
        wb.close()


def read_rows(data: bytes, filename: str, mapping: Mapping, max_rows: int = MAX_ROWS) -> list[dict]:
    suffix = Path(filename).suffix.lower()
    if suffix not in (".csv", ".xlsx"):
        raise ValueError(f"unsupported file type: {suffix or filename!r}")
    rows = []
    try:
        for n, rec in _records(data, suffix, mapping, max_rows):
            rec = {k.strip(): _cell(v) for k, v in rec.items() if isinstance(k, str)}
            if all(v is None or v == "" for v in rec.values()):
                continue
            if len(rows) >= max_rows:
                raise ValueError(f"too many rows (more than {max_rows})")
            row = {c: rec[h] for c, h in mapping.columns.items() if h in rec}
            row["_row"] = n
            rows.append(row)
    except ValueError:
        raise
    except Exception as exc:  # untrusted input: zip, csv, openpyxl and sheet-lookup errors all surface as one type
        raise ValueError(f"unreadable file: {type(exc).__name__}: {str(exc)[:100]}") from exc
    return rows


def template(kind: str, fmt: Literal["csv", "xlsx"]) -> bytes:
    cols = load_mapping("default").columns
    header = [cols[c] for c in TEMPLATE_COLUMNS[kind]]
    if fmt == "csv":
        buf = io.StringIO()
        csv.writer(buf, lineterminator="\n").writerow(header)
        return buf.getvalue().encode()
    wb = openpyxl.Workbook()
    wb.active.append(header)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def file_ref(data: bytes) -> str:
    return "sha256:" + hashlib.sha256(data).hexdigest()
