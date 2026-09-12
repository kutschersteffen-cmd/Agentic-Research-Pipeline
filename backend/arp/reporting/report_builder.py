from __future__ import annotations

import tempfile
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt

from arp.reporting import chart_builder
from arp.schemas.reporting import LayoutInstructions, QuantitativeDataset, ReportPlan, ReportSection, SectionLayoutHint

_IMAGE_WIDTH_IN = 6.0


def _add_narrative(document: Document, section: ReportSection) -> None:
    for item in section.narrative:
        style = "List Bullet" if item.bullet else "Normal"
        document.add_paragraph(item.text, style=style)


def _add_table(document: Document, table_spec, datasets: list[QuantitativeDataset]) -> None:
    ds = next((d for d in datasets if d.dataset_id == table_spec.dataset_id), None)
    if ds is None:
        return
    columns = table_spec.columns or ds.column_names()
    rows = ds.rows[: table_spec.max_rows]
    table = document.add_table(rows=len(rows) + 1, cols=len(columns))
    table.style = "Light Grid Accent 1"
    for c, name in enumerate(columns):
        cell = table.cell(0, c)
        cell.text = name
        cell.paragraphs[0].runs[0].bold = True
    for r, row in enumerate(rows, start=1):
        for c, name in enumerate(columns):
            table.cell(r, c).text = "" if row.get(name) is None else str(row.get(name))
    if len(ds.rows) > table_spec.max_rows:
        note = document.add_paragraph(f"(+{len(ds.rows) - table_spec.max_rows} more rows in the underlying dataset)")
        note.runs[0].italic = True
        note.runs[0].font.size = Pt(9)


def _add_chart(document: Document, section: ReportSection, datasets: list[QuantitativeDataset], tmp_dir: Path) -> None:
    spec = section.chart
    png_path = tmp_dir / f"chart_{id(spec)}.png"
    chart_builder.render_chart_image(spec, datasets, png_path)
    document.add_picture(str(png_path), width=Inches(_IMAGE_WIDTH_IN))
    if spec.notes:
        caption = document.add_paragraph(spec.notes)
        caption.runs[0].italic = True
        caption.runs[0].font.size = Pt(9)


def _ordered_sections(plan: ReportPlan, layout: LayoutInstructions) -> list[ReportSection]:
    if not layout.include_appendix:
        return list(plan.sections)
    main = [s for s in plan.sections if not s.appendix]
    appendix = [s for s in plan.sections if s.appendix]
    return main + appendix


def build_docx(plan: ReportPlan, datasets: list[QuantitativeDataset], layout: LayoutInstructions, out_path: Path) -> Path:
    """Deterministically renders a ReportPlan into a .docx report: one
    heading + body per section, flowing continuously (no slide/page
    concept the way DeckBuilder has one). Charts have no native, editable
    equivalent in the docx object model, so they're always embedded as
    static images (see chart_builder.render_chart_image) -- unlike the
    pptx path, which keeps most chart types as live Office chart objects.
    """
    document = Document()
    document.add_heading(plan.title, level=0)
    if plan.subtitle:
        sub = document.add_paragraph(plan.subtitle)
        sub.runs[0].italic = True

    with tempfile.TemporaryDirectory(prefix="arp_report_chart_") as tmp:
        tmp_dir = Path(tmp)
        sections = _ordered_sections(plan, layout)
        appendix_started = False
        for section in sections:
            if layout.include_appendix and section.appendix and not appendix_started:
                appendix_started = True
                document.add_heading("Appendix", level=1)

            level = 2 if section.layout_hint == SectionLayoutHint.SECTION_HEADER else 1
            document.add_heading(section.heading, level=level)
            if section.layout_hint == SectionLayoutHint.SECTION_HEADER:
                continue

            _add_narrative(document, section)
            if section.chart is not None:
                _add_chart(document, section, datasets, tmp_dir)
            elif section.table is not None:
                _add_table(document, section.table, datasets)

        out_path.parent.mkdir(parents=True, exist_ok=True)
        document.save(str(out_path))
    return out_path
