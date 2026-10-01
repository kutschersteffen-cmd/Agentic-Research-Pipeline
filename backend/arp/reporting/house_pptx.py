"""Deck -> editable .pptx from the same tokens and slot grid as the HTML renderer.

Shapes sit on the blank layout at slot_rect positions (named `slot:<name>`), so the
positions match the HTML by construction. Text is the already-fitted deck text; nothing
is re-fitted here. Fonts are named, not embedded (python-pptx cannot embed them).
"""

from __future__ import annotations

import math
import tempfile
import textwrap
from pathlib import Path
from typing import NamedTuple

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from arp.reporting.chart_builder import build_native_chart_data, is_native, render_chart_image, style_native_chart
from arp.reporting.house_style import Mode, Rect, Tokens, TypeStyle, get_variant, load_tokens, slot_rect
from arp.reporting.html_render import _DISPLAY_LAYOUTS, Density, _heat, table_view, theme_from_tokens
from arp.reporting.structured import _BOX_H, parse_tree, structured_view, tree_layout
from arp.schemas.reporting import Deck, QuantitativeDataset, SlideContent

_PX = 9525  # EMU per px
_ROW_PX = 46  # table row: caption line + 2x10px padding, as in the HTML


def _e(px: float) -> int:
    return round(px * _PX)


def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color.lstrip("#"))


def _family(stack: str) -> str:
    return stack.split(",")[0].strip().strip("'\"")


def _srgb(color: str) -> tuple[float, float, float]:
    """'#rrggbb' or the tokens' 'oklch(L% C h)' as 0-1 sRGB."""
    if not color.startswith("oklch"):
        return tuple(int(color.lstrip("#")[i : i + 2], 16) / 255 for i in (0, 2, 4))
    lt, c, h = (float(v.rstrip("%")) for v in color[color.index("(") + 1 : -1].split())
    lt, a, b = lt / 100, c * math.cos(math.radians(h)), c * math.sin(math.radians(h))
    l_, m_, s_ = ((lt + k1 * a + k2 * b) ** 3 for k1, k2 in ((0.3963377774, 0.2158037573), (-0.1055613458, -0.0638541728), (-0.0894841775, -1.2914855480)))
    lin = (4.0767416621 * l_ - 3.3077115913 * m_ + 0.2309699292 * s_, -1.2684380046 * l_ + 2.6097574011 * m_ - 0.3413193965 * s_,
           -0.0041960863 * l_ - 0.7034186147 * m_ + 1.7076147010 * s_)
    return tuple(min(1, max(0, 12.92 * v if v <= 0.0031308 else 1.055 * v ** (1 / 2.4) - 0.055)) for v in lin)


class Para(NamedTuple):
    """One paragraph; fields left None take the shape's defaults. `text` may be [(run, bold)] for mixed weight."""

    text: str | list[tuple[str, bool]]
    ts: TypeStyle | None = None
    font: str | None = None
    color: str | None = None
    before: float = 0
    after: float | None = None
    bullet: str | None = None


class _Builder:
    def __init__(self, tokens: Tokens, mode: Mode, density: Density = "committee"):
        self.t, self.mode, self.c, self.density = tokens, mode, tokens.colors(mode), density
        self.theme = theme_from_tokens(tokens, mode)
        self.body, self.mono = _family(tokens.fonts.body), _family(tokens.fonts.mono or tokens.fonts.body)
        self.head = _family(tokens.heading_font(mode))

    def _tint(self, status: str, pct: float) -> str:
        """The status colour mixed into the ground, as the HTML's color-mix does (here in sRGB)."""
        mix = [pct * s + (1 - pct) * b for s, b in zip(_srgb(self.c.status[status]), _srgb(self.c.background), strict=True)]
        return "#" + "".join(f"{round(v * 255):02x}" for v in mix)

    def _est_h(self, text: str, ts: TypeStyle, w: float) -> float:
        """Wrapped text height; ponytail: average-advance estimate (as structured.py's), a measured wrap if a face runs wide."""
        return max(1, len(textwrap.wrap(text, max(1, int(w / (ts.size * 0.52)))))) * ts.size * ts.line_height

    def _role(self, spec, layout: str, value) -> TypeStyle:
        """The HTML's role choice: committee prose sets at body size; otherwise short content steps up to type_role_short."""
        if self.density == "committee" and layout not in _DISPLAY_LAYOUTS and spec.kind in ("text", "list"):
            return self.t.type["body" if spec.type_role in ("subhead", "headline") else spec.type_role]
        words = max((len(x.split()) for x in value), default=0) if isinstance(value, list) else len((value or "").split())
        return self.t.type[spec.type_role_short if spec.type_role_short and 0 < words <= (spec.short_words or 0) else spec.type_role]

    def _line(self, slide, name: str, p1: tuple[float, float], p2: tuple[float, float], color: str, px: float, *, elbow=False, arrow=False) -> None:
        ln = slide.shapes.add_connector(MSO_CONNECTOR.ELBOW if elbow else MSO_CONNECTOR.STRAIGHT, _e(p1[0]), _e(p1[1]), _e(p2[0]), _e(p2[1]))
        ln.name = name
        ln.line.color.rgb, ln.line.width = _rgb(color), Emu(_e(px))
        if arrow:
            etree.SubElement(ln.line._get_or_add_ln(), qn("a:tailEnd")).set("type", "triangle")

    def _box(self, slide, name: str, r: Rect, shape=None):
        shape = slide.shapes.add_shape(shape, _e(r.x), _e(r.y), _e(r.w), _e(r.h)) if shape else slide.shapes.add_textbox(_e(r.x), _e(r.y), _e(r.w), _e(r.h))
        shape.name = name
        # Float rects round to the same EMU as the HTML px, but add_* may truncate; pin them.
        shape.left, shape.top, shape.width, shape.height = _e(r.x), _e(r.y), _e(r.w), _e(r.h)
        return shape

    def _fill(self, shape, hex_color: str) -> None:
        shape.fill.solid()
        shape.fill.fore_color.rgb = _rgb(hex_color)
        shape.line.fill.background()
        shape.shadow.inherit = False

    def _rule(self, slide, name: str, r: Rect, color: str) -> None:
        self._fill(self._box(slide, name, r, MSO_SHAPE.RECTANGLE), color)

    def _text(self, slide, name: str, r: Rect, paras: list[str | Para], ts: TypeStyle, *, font: str, color: str, wrap=True, anchor=MSO_ANCHOR.TOP,
              align=PP_ALIGN.LEFT, bullet: str | None = None, pad_top: float = 0, ml: float = 0, mr: float = 0, mb: float = 0,
              fill: str | None = None, line: tuple[str, float] | None = None, shape=None, vert: str | None = None):
        shape = self._box(slide, name, r, shape)
        if fill:
            shape.fill.solid()
            shape.fill.fore_color.rgb = _rgb(fill)
        if line:
            shape.line.color.rgb, shape.line.width = _rgb(line[0]), Emu(_e(line[1]))
        elif shape.shape_type != 17:  # autoshapes get the theme's outline unless told otherwise; text boxes have none
            shape.line.fill.background()
        tf = shape.text_frame
        tf.word_wrap, tf.vertical_anchor, tf.auto_size = wrap, anchor, MSO_AUTO_SIZE.NONE
        tf.margin_left, tf.margin_right, tf.margin_bottom, tf.margin_top = _e(ml), _e(mr), _e(mb), _e(pad_top)
        if vert:
            tf._txBody.bodyPr.set("vert", vert)
        for i, para in enumerate(paras):
            para = para if isinstance(para, Para) else Para(para)
            pts, bu = para.ts or ts, para.bullet or bullet
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment, p.line_spacing = align, Pt(pts.size * pts.line_height * 0.75)
            if para.after is not None:
                p.space_after = Pt(para.after * 0.75)
            elif bu and len(paras) > 1 or bu == "num":
                p.space_after = Pt(pts.size * 0.5 * 0.75)
            if para.before:
                p.space_before = Pt(para.before * 0.75)
            for text, bold in para.text if isinstance(para.text, list) else [(para.text, pts.weight >= 600)]:
                run = p.add_run()
                run.text = text
                f = run.font
                f.name, f.size, f.bold, f.color.rgb = para.font or font, Pt(pts.size * 0.75), bold, _rgb(para.color or color)
                if pts.tracking:
                    run._r.get_or_add_rPr().set("spc", str(round(pts.tracking * pts.size * 0.75 * 100)))
            if bu:
                pPr = p._p.get_or_add_pPr()
                pPr.set("marL", str(_e(pts.size * 1.1 if bu != "num" else pts.size * 1.5)))
                pPr.set("indent", str(-_e(pts.size * 1.1 if bu != "num" else pts.size * 1.5)))
                el = etree.SubElement(pPr, qn("a:buAutoNum" if bu == "num" else "a:buChar"))
                el.set("type", "arabicPeriod") if bu == "num" else el.set("char", "•")
        return shape

    def _table(self, slide, spec, r: Rect, slide_content: SlideContent, ds: QuantitativeDataset) -> None:
        columns, rows, more = table_view(slide_content.table, ds)
        if not columns:
            return
        ts, mono = self.t.type[spec.type_role], _family(self.t.fonts.mono or self.t.fonts.body)
        cap = self.t.type["caption"]
        heat = _heat(slide_content.table, columns, rows) if slide_content.table.heat else None
        n = len(rows) + 1
        shape = slide.shapes.add_table(n, max(len(columns), 1), _e(r.x), _e(r.y), _e(r.w), _e(_ROW_PX * n))
        shape.name = f"slot:{spec.name}"
        tbl = shape.table
        tbl.horz_banding = False
        for row in tbl.rows:
            row.height = _e(_ROW_PX)
        for i, col in enumerate(tbl.columns):
            col.width = _e(r.w / len(columns)) if i < len(columns) - 1 else _e(r.w) - _e(r.w / len(columns)) * (len(columns) - 1)
        for ri, values in enumerate([columns, *rows]):
            for ci, value in enumerate(values):
                cell = tbl.cell(ri, ci)
                cell.fill.solid()
                cell.fill.fore_color.rgb = _rgb(self._tint(heat[ri - 1][ci], 0.16) if ri and heat and heat[ri - 1][ci] else self.c.background)
                cell.margin_left = cell.margin_right = _e(12)
                cell.margin_top = cell.margin_bottom = _e(10)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                run = cell.text_frame.paragraphs[0].add_run()
                f = run.font
                if ri == 0:  # small uppercase mono header, as deck.css.j2
                    run.text = str(value).upper()
                    f.name, f.size, f.bold = mono, Pt(cap.size * 0.75), False
                    run._r.get_or_add_rPr().set("spc", str(round(0.08 * cap.size * 0.75 * 100)))
                else:
                    run.text = str(value)
                    f.name, f.size, f.bold = _family(self.t.fonts.body), Pt(ts.size * 0.75), False
                f.color.rgb = _rgb(self.c.ink_muted if ri == 0 else self.c.ink)
                tcPr = cell._tc.get_or_add_tcPr()
                # a:tcPr takes lnL, lnR, lnT, lnB (then the fill); the table style's default grid is switched off on three sides.
                for i, side in enumerate(("lnL", "lnR", "lnT", "lnB")):
                    ln = etree.Element(qn(f"a:{side}"), w=str(_e(2 if ri == 0 else 1)))
                    if side == "lnB":
                        etree.SubElement(etree.SubElement(ln, qn("a:solidFill")), qn("a:srgbClr")).set("val", (self.c.ink if ri == 0 else self.c.neutral).lstrip("#"))
                    else:
                        etree.SubElement(ln, qn("a:noFill"))
                    tcPr.insert(i, ln)
        shape.left, shape.top, shape.width, shape.height = _e(r.x), _e(r.y), _e(r.w), _e(r.h)
        if more:
            y = min(r.y + _ROW_PX * n + 8, r.y + r.h - ts.size * ts.line_height)
            self._text(slide, f"slot:{spec.name}:more", Rect(r.x, y, r.w, ts.size * ts.line_height), [f"+{more} more rows"], ts,
                       font=_family(self.t.fonts.body), color=self.c.ink_muted)

    def _picture(self, slide, spec, r: Rect, path: Path) -> None:
        if not path.is_file():
            raise FileNotFoundError(f"image_path does not exist: {path}")
        pic = slide.shapes.add_picture(str(path), _e(r.x), _e(r.y), _e(r.w), _e(r.h))
        pic.name = f"slot:{spec.name}"
        with Image.open(path) as im:
            w, h = im.size
        # object-fit: contain, keeping the frame at slot_rect: negative crop pads the short side.
        if w / h > r.w / r.h:
            pic.crop_top = pic.crop_bottom = -((r.h * w / r.w / h) - 1) / 2
        else:
            pic.crop_left = pic.crop_right = -((r.w * h / r.h / w) - 1) / 2

    # ---- committee patterns: every composite keeps an empty `slot:<name>` frame at slot_rect, its parts are named `slot:<name>:...`

    def _frame(self, slide, name: str, r: Rect) -> None:
        self._box(slide, f"slot:{name}", r)

    def _cards_items(self, slide, spec, content, items, r, ts):
        g, sub, n = self.t.grid.gutter, self.t.type["subhead"], len(items)
        w = (r.w - (n - 1) * g) / n
        cards = [i.split(": ", 1) for i in items]
        h = max(2 * g + sub.size + 20 + (self._est_h(c[0], sub, w - 2 * g) if len(c) > 1 else 0) + (12 if len(c) > 1 else 0)
                + self._est_h(c[-1], ts, w - 2 * g) + 2 for c in cards)
        y = r.y if self.density == "committee" else max(r.y, r.y + (r.h - h) / 2)  # a pitch centres the row; a pre-read reads from the top
        for i, c in enumerate(cards):
            paras = [Para(f"{i + 1:02d}", sub.model_copy(update={"weight": 500, "line_height": 1}), self.mono, self.c.ink_muted, after=20)]
            if len(c) > 1:
                paras.append(Para(c[0], sub, self.head, self.c.ink, after=0))
            paras.append(Para(c[-1], ts.model_copy(update={"weight": 400}), self.body, self.c.ink_muted if len(c) > 1 else self.c.ink, before=12 if len(c) > 1 else 0))
            self._text(slide, f"slot:items:{i}", Rect(r.x + i * (w + g), y, w, h), paras, ts, font=self.body, color=self.c.ink,
                       ml=g, mr=g, pad_top=g, line=(self.c.neutral, 1))

    def _steps_items(self, slide, spec, content, items, r, ts):
        g, sub, body, n = self.t.grid.gutter, self.t.type["subhead"], self.t.type["body"], len(items)
        w = (r.w - (n - 1) * g) / n
        steps = [i.split(": ", 1) for i in items]
        h = 112 + max((self._est_h(c[0], ts, w) + 12 if len(c) > 1 else 0) + self._est_h(c[-1], body if len(c) > 1 else ts, w) for c in steps)
        y = max(r.y, r.y + (r.h - h) / 2)
        self._line(slide, "slot:items:line", (r.x, y + 36), (r.x + r.w, y + 36), self.c.ink_muted, 2)
        for i, c in enumerate(steps):
            x = r.x + i * (w + g)
            self._text(slide, f"slot:items:{i}:node", Rect(x, y, 72, 72), [str(i + 1)], body.model_copy(update={"weight": 500}), font=self.mono, color=self.c.ink,
                       anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER, fill=self.c.background, line=(self.c.ink, 2), shape=MSO_SHAPE.OVAL)
            paras = ([Para(c[0], ts.model_copy(update={"weight": sub.weight}), self.head, self.c.ink)] if len(c) > 1 else [])
            paras.append(Para(c[-1], body if len(c) > 1 else ts, self.body, self.c.ink_muted if len(c) > 1 else self.c.ink, before=12 if len(c) > 1 else 0))
            self._text(slide, f"slot:items:{i}", Rect(x, y + 112, w, h - 112), paras, ts, font=self.body, color=self.c.ink)

    def _rows_items(self, slide, spec, content, items, r, ts):
        """summary / split list: full-width rows on hairlines sharing the height, a mono index in a fixed column."""
        h, body = r.h / len(items), self.t.type["body"]
        for i, text in enumerate(items):
            y = r.y + i * h
            self._rule(slide, f"slot:items:{i}:rule", Rect(r.x, y, r.w, 1), self.c.neutral)
            self._text(slide, f"slot:items:{i}:index", Rect(r.x, y, 112, h), [f"{i + 1:02d}"], body.model_copy(update={"weight": 500}), font=self.mono,
                       color=self.c.ink_muted, anchor=MSO_ANCHOR.MIDDLE)
            self._text(slide, f"slot:items:{i}", Rect(r.x + 112, y, r.w - 112, h), [text], ts, font=self.body, color=self.c.ink, anchor=MSO_ANCHOR.MIDDLE)
        self._rule(slide, "slot:items:rule", Rect(r.x, r.y + r.h - 1, r.w, 1), self.c.neutral)

    def _flow_items(self, slide, spec, content, items, r, ts):
        st = structured_view("flow", "items", items, r)["stages"]
        n, col_h, body = len(st), r.h - 24, self.t.type["body"]
        w = (r.w - (n - 1) * 48) / n
        badge_h, badge_w = body.size * self.t.type["caption"].line_height, 9 * body.size * 0.72 + 16
        for i, stage in enumerate(st):
            x = r.x + i * (w + 48)
            self._text(slide, f"slot:items:{i}", Rect(x, r.y, w, col_h), [Para(f"{i + 1} {stage.title}", ts.model_copy(update={"weight": 600}))], ts, font=self.body,
                       color=self.c.ink, ml=20, mr=20, pad_top=20, fill=self.c.well)
            hs = [self._est_h(t, ts, w - 72) + 24 + (badge_h + 8 if pref else 0) for t, pref in stage.boxes]
            title_h = self._est_h(f"{i + 1} {stage.title}", ts, w - 40)
            y = max(r.y + 20 + title_h + 12, r.y + (col_h - sum(hs) - 12 * (len(hs) - 1)) / 2)
            for j, ((text, pref), h) in enumerate(zip(stage.boxes, hs, strict=True)):
                self._text(slide, f"slot:items:{i}:box:{j}", Rect(x + 20, y, w - 40, h), [text], ts.model_copy(update={"weight": 400}), font=self.body, color=self.c.ink,
                           ml=16, mr=16, pad_top=12 + (badge_h + 8 if pref else 0), fill=self.c.background, line=(self.c.ink, 2) if pref else (self.c.neutral, 1))
                if pref:
                    self._text(slide, f"slot:items:{i}:box:{j}:badge", Rect(x + 36, y + 12, badge_w, badge_h), ["PREFERRED"],
                               body.model_copy(update={"weight": 500, "tracking": 0.06}), font=self.mono, color=self.c.background, wrap=False, ml=8, fill=self.c.ink)
                y += h + 12
            if i < n - 1:
                self._line(slide, f"slot:items:arrow:{i}", (x + w + 12, r.y + col_h / 2), (x + w + 36, r.y + col_h / 2), self.c.ink_muted, 2, arrow=True)

    def _meters(self, slide, spec, content, items, r, ts):
        y = r.y
        for i, m in enumerate(structured_view("profile", "meters", items, r)["meters"]):
            h = ts.size * ts.line_height
            self._text(slide, f"slot:meters:{i}", Rect(r.x, y, r.w - 160, h), [m.label], ts, font=self.body, color=self.c.ink)
            self._text(slide, f"slot:meters:{i}:pct", Rect(r.x + r.w - 160, y, 160, h), [f"{round(m.pct)}%"], ts.model_copy(update={"weight": 500}), font=self.mono,
                       color=self.c.ink, align=PP_ALIGN.RIGHT)
            self._rule(slide, f"slot:meters:{i}:track", Rect(r.x, y + h + 8, r.w, 10), self.c.neutral)
            self._rule(slide, f"slot:meters:{i}:fill", Rect(r.x, y + h + 8, r.w * m.pct / 100, 10), self.c.ink)
            y += h + 8 + 10 + 14

    def _panel(self, slide, spec, content, items, r, ts):
        v = structured_view(content.layout, spec.name, items, r)
        sub = self.t.type["subhead"]
        paras = [Para(v["panel"][0], sub, self.head, self.c.ink, after=12)] + [Para(t, ts, self.body, self.c.ink, bullet="dot", after=0.4 * ts.size) for t in v["items"]]
        self._text(slide, f"slot:{spec.name}:panel", Rect(r.x, r.y, r.w, r.h - 24), paras, ts, font=self.body, color=self.c.ink, ml=24, mr=24, pad_top=20, line=(self.c.neutral, 1))

    def _deltas(self, slide, spec, content, items, r, ts):
        head, y = self.t.type["headline"], r.y
        for i, d in enumerate(structured_view("scatter_zone", "items", items, r)["deltas"]):
            h = max(head.size, self._est_h(d.text, ts, r.w - 204))
            self._text(slide, f"slot:items:{i}:value", Rect(r.x, y, 180, head.size), [d.value], head.model_copy(update={"line_height": 1}), font=self.head, color=self.c.ink, wrap=False)
            self._text(slide, f"slot:items:{i}", Rect(r.x + 204, y, r.w - 204, h), [d.text], ts, font=self.body, color=self.c.ink)
            y += h + 32

    def _quadrants(self, slide, spec, content, items, r, ts):
        w, h = (r.w - 16) / 2, (r.h - 24 - 16) / 2
        for i, q in enumerate(structured_view("matrix2x2", "quadrants", items, r)["quadrants"]):
            paras = [Para(q.title, ts.model_copy(update={"weight": 600})), Para(q.sub, ts.model_copy(update={"weight": 400}), color=self.c.ink_muted, after=8),
                     Para(q.text, ts.model_copy(update={"weight": 400}))]
            self._text(slide, f"slot:quadrants:{i}", Rect(r.x + (i % 2) * (w + 16), r.y + (i // 2) * (h + 16), w, h), paras, ts, font=self.body, color=self.c.ink,
                       ml=20, mr=20, pad_top=16, fill=self._tint(q.status, 0.16))

    def _tree(self, slide, spec, content, items, r, ts):
        box_w, nodes, edges = tree_layout(parse_tree(items), r.w, r.h)
        t24 = TypeStyle(size=24, line_height=1.25, weight=400, font="body")  # the SVG's 24px floor
        for i, (n, x, y, lines) in enumerate(nodes):
            self._text(slide, f"slot:items:node:{i}", Rect(r.x + x, r.y + y - _BOX_H / 2, box_w, _BOX_H), lines,
                       t24.model_copy(update={"weight": 400 if n.outcome else 600}), font=self.body, color=self.c.ink, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER,
                       ml=12, mr=12, fill=self._tint(n.status, 0.18) if n.outcome else self.c.background,
                       line=(self.c.neutral, 1) if n.outcome else (self.c.ink, 2))
        for k, (label, x0, y, xm, cy, xe) in enumerate(edges):
            self._line(slide, f"slot:items:edge:{k}", (r.x + x0, r.y + y), (r.x + xe, r.y + cy), self.c.ink_muted, 2, elbow=True, arrow=True)
            self._text(slide, f"slot:items:edge:{k}:label", Rect(r.x + xm + 8, r.y + cy - 34, 80, 30), [label], t24.model_copy(update={"weight": 500}), font=self.mono,
                       color=self.c.ink_muted, wrap=False)

    def _decisions(self, slide, spec, content, items, r, ts):
        sub, head, y = self.t.type["subhead"], self.t.type["headline"], r.y
        for i, (title, detail) in enumerate(structured_view("decisions", "items", items, r)["pairs"]):
            h = 32 + sub.size * sub.line_height + (self._est_h(detail, ts, r.w - 192) if detail else 0)
            paras = [Para(title, sub, self.head, self.c.ink)] + ([Para(detail, ts, self.body, self.c.ink_muted)] if detail else [])
            self._text(slide, f"slot:items:{i}", Rect(r.x, y, r.w, h), paras, ts, font=self.body, color=self.c.ink, anchor=MSO_ANCHOR.MIDDLE, ml=128, mr=32, line=(self.c.neutral, 1))
            self._text(slide, f"slot:items:{i}:num", Rect(r.x + 32, y, 96, h), [str(i + 1)], head.model_copy(update={"line_height": 1}), font=self.head, color=self.c.ink,
                       anchor=MSO_ANCHOR.MIDDLE)
            y += h + 12

    def _compare(self, slide, spec, r, text, ts):
        """Two mirrored panels under an ink rule; the "Label:" before the first colon heads its panel."""
        g, p = self.t.grid.gutter, text.split(": ", 1)
        head = self.t.type["subhead" if self.density == "committee" else "headline"]
        self._rule(slide, f"slot:{spec.name}:rule", Rect(r.x, r.y, r.w, 2), self.c.ink)
        paras = ([Para(p[0], head.model_copy(update={"tracking": 0}) if self.density == "committee" else head, self.head, self.c.ink, after=16)] if len(p) > 1 else [])
        paras.append(Para(p[-1], ts.model_copy(update={"weight": 400}), self.body, self.c.ink))
        self._text(slide, f"slot:{spec.name}", r, paras, ts, font=self.body, color=self.c.ink, pad_top=g + 2)

    _COMPOSITE = {("cards", "items"): _cards_items, ("steps", "items"): _steps_items, ("summary", "items"): _rows_items, ("split", "items"): _rows_items,
                  ("flow", "items"): _flow_items, ("profile", "meters"): _meters, ("profile", "left"): _panel, ("profile", "right"): _panel,
                  ("matrix2x2", "items"): _panel, ("matrix2x2", "quadrants"): _quadrants, ("scatter_zone", "items"): _deltas, ("tree", "items"): _tree,
                  ("decisions", "items"): _decisions}

    def _slot(self, slide, i: int, spec, content: SlideContent, datasets: dict[str, QuantitativeDataset]) -> None:
        r, value = slot_rect(self.t, spec), content.slots.get(spec.name)
        ts = self._role(spec, content.layout, value)
        name, body = f"slot:{spec.name}", _family(self.t.fonts.body if ts.font == "body" else self.t.heading_font(self.mode))
        if spec.kind == "chart" and content.chart:
            if is_native(content.chart) and content.layout != "scatter_zone":  # a native scatter cannot shade a zone; that one is the picture
                kind, data = build_native_chart_data(content.chart, list(datasets.values()))
                frame = slide.shapes.add_chart(kind, _e(r.x), _e(r.y), _e(r.w), _e(r.h), data)
                frame.name = name
                style_native_chart(frame.chart, content.chart, self.theme, min_pt=18)  # 24px floor on the 1920px canvas
            else:
                with tempfile.TemporaryDirectory() as tmp:
                    img = render_chart_image(content.chart, list(datasets.values()), Path(tmp) / "c.png", width_in=r.w / 96, height_in=r.h / 96, theme=self.theme)
                    slide.shapes.add_picture(str(img), _e(r.x), _e(r.y), _e(r.w), _e(r.h)).name = name
        elif spec.kind == "table" and content.table and content.table.dataset_id in datasets:
            self._table(slide, spec, r, content, datasets[content.table.dataset_id])
        elif spec.kind == "image" and content.image_path:
            self._picture(slide, spec, r, Path(content.image_path))
        elif spec.kind == "list":
            items = value if isinstance(value, list) else ([value] if value else [])
            if items and (fn := self._COMPOSITE.get((content.layout, spec.name))):
                try:
                    self._frame(slide, spec.name, r)
                    return fn(self, slide, spec, content, items, r, ts)
                except ValueError:  # malformed structure: a plain list, as the HTML does (its frame is replaced below)
                    for sh in [s for s in slide.shapes if s.name == name or s.name.startswith(name + ":")]:
                        sh._element.getparent().remove(sh._element)
            # Timeline: numbered steps; other lists get bullets.
            self._text(slide, name, r, items, ts, font=body, color=self.c.ink, bullet="num" if content.layout == "timeline" else "dot")
        else:
            text = " ".join(value) if isinstance(value, list) else (value or "")
            lay, mute = content.layout, self.c.ink_muted
            if spec.kind == "number":
                # Metric numbers are ink (DESIGN.md); bottom-aligned on its label, like the HTML; the descent padding is 0.2em
                self._text(slide, name, Rect(r.x, r.y, r.w, r.h), [text], ts, font=body, color=self.c.ink, wrap=False, anchor=MSO_ANCHOR.BOTTOM)
            elif spec.name.startswith("label_"):
                self._rule(slide, f"{name}:rule", Rect(r.x, r.y, r.w, 1), self.c.neutral)
                self._text(slide, name, r, [text], ts, font=self.mono, color=self.c.ink, pad_top=16)
            elif lay == "compare" and spec.name in ("left", "right"):
                self._compare(slide, spec, r, text, ts)
            elif spec.name == "takeaway_bar" and text:
                self._frame(slide, spec.name, r)
                p = text.split(": ", 1)
                h = self._est_h(text, ts, r.w - 48) + 28
                self._text(slide, f"{name}:bar", Rect(r.x, r.y + r.h - h, r.w, h), [Para([(p[0] + ":", True), (" " + p[1], False)] if len(p) > 1 else p[0])], ts,
                           font=body, color=self.c.ink, anchor=MSO_ANCHOR.MIDDLE, ml=24, mr=24, fill=self.c.well)
            elif spec.name in ("x_axis", "y_axis"):  # mono labels; the y axis reads bottom to top
                self._text(slide, name, r, [text.upper()], ts.model_copy(update={"weight": 500, "tracking": 0.06}), font=self.mono, color=mute, align=PP_ALIGN.CENTER,
                           anchor=MSO_ANCHOR.MIDDLE if spec.name == "y_axis" else MSO_ANCHOR.TOP, vert="vert270" if spec.name == "y_axis" else None)
            elif lay in ("title", "section") and spec.name == "subtitle":
                self._rule(slide, f"{name}:rule", Rect(r.x, r.y, r.w, 1), self.c.neutral)
                self._text(slide, name, r, [text], ts, font=body, color=mute, pad_top=self.t.grid.gutter)
            elif (lay == "section" and spec.name == "number") or spec.name == "attribution":
                self._text(slide, name, r, [("— " if spec.name == "attribution" else "") + text], ts.model_copy(update={"weight": 500, "tracking": 0}), font=self.mono, color=mute,
                           pad_top=self.t.grid.gutter if spec.name == "attribution" else 0)
            else:
                anchor = (MSO_ANCHOR.BOTTOM if (lay in ("title", "section") and spec.name == "title") or lay == "quote" and spec.name == "quote"
                          or (lay, content.variant, spec.name) == ("statement", "support", "statement")
                          else MSO_ANCHOR.MIDDLE if (lay, content.variant, spec.name) == ("statement", "plain", "statement") else MSO_ANCHOR.TOP)
                self._text(slide, name, r, [text] if text else [], ts, font=body, color=mute if (lay, spec.name) == ("statement", "support") else self.c.ink, anchor=anchor,
                           pad_top=self.t.grid.gutter if (lay, spec.name) == ("statement", "support") else 0, mb=0.2 * ts.size if anchor == MSO_ANCHOR.BOTTOM else 0)

    def slide(self, prs, i: int, content: SlideContent, datasets: dict[str, QuantitativeDataset]) -> None:
        t, g = self.t, self.t.grid
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = _rgb(self.c.background)
        self._rule(slide, "accent_rule", Rect(g.margin_x, g.margin_top - 32 - g.rule_height, g.rule_width, g.rule_height), self.c.ink)
        if content.layout not in ("title", "section"):  # those show the title in their own slot, as in the HTML
            self._text(slide, "headline", Rect(g.margin_x, g.margin_top, t.canvas.width - 2 * g.margin_x, g.headline_band - 32), [content.headline], t.type["headline"],
                       font=_family(t.heading_font(self.mode)), color=self.c.ink)
        if content.eyebrow:
            cap = t.type["caption"]
            self._text(slide, "eyebrow", Rect(g.margin_x + g.rule_width + 16, g.margin_top - 32 - g.rule_height // 2 - 18, t.canvas.width - 2 * g.margin_x - g.rule_width - 16, 36),
                       [content.eyebrow.upper()], t.type["body"].model_copy(update={"weight": 500, "tracking": 0.06, "line_height": cap.line_height}), font=self.mono,
                       color=self.c.ink_muted, wrap=False, anchor=MSO_ANCHOR.MIDDLE)
        for spec in get_variant(content.layout, content.variant).slots:
            self._slot(slide, i, spec, content, datasets)
        top, mono = t.canvas.height - g.margin_bottom - g.footer_band, _family(t.fonts.mono or t.fonts.body)
        cap = t.type["caption"]
        self._rule(slide, "footer:rule", Rect(g.margin_x, top, t.canvas.width - 2 * g.margin_x, 1), self.c.neutral)
        row = Rect(g.margin_x, top, t.canvas.width - 2 * g.margin_x, g.footer_band)
        self._text(slide, "footer:refs", row, [" · ".join(content.source_refs)], cap, font=mono, color=self.c.ink_muted, anchor=MSO_ANCHOR.MIDDLE)
        self._text(slide, "footer:page", row, [str(i + 1)], cap, font=mono, color=self.c.ink_muted, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.RIGHT)
        if content.speaker_notes:
            slide.notes_slide.notes_text_frame.text = content.speaker_notes


def build_house_pptx(deck: Deck, datasets: list[QuantitativeDataset], out: Path, tokens: Tokens | None = None, mode: Mode = "light",
                    density: Density = "committee") -> Path:
    tokens = tokens or load_tokens()
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(_e(tokens.canvas.width)), Emu(_e(tokens.canvas.height))
    b, by_id = _Builder(tokens, mode, density), {d.dataset_id: d for d in datasets}
    for i, content in enumerate(deck.slides):
        b.slide(prs, i, content, by_id)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    return out
