import asyncio

from docx import Document

from arp.reporting.browser import write_pdf
from arp.reporting.deck_compat import deck_from_plan
from arp.reporting.html_render import render_deck_html
from arp.reporting.report_builder import build_docx
from arp.schemas.reporting import (
    ChartSpec,
    ChartType,
    ColumnKind,
    ContentItem,
    DatasetColumn,
    LayoutInstructions,
    QuantitativeDataset,
    ReportPlan,
    ReportSection,
    SectionLayoutHint,
    TableSpec,
)


def _dataset() -> QuantitativeDataset:
    return QuantitativeDataset(
        name="Revenue",
        columns=[DatasetColumn(name="segment", kind=ColumnKind.CATEGORY), DatasetColumn(name="revenue", kind=ColumnKind.NUMBER)],
        rows=[{"segment": "EV", "revenue": 120}, {"segment": "Grid", "revenue": 80}],
    )


def _plan(ds: QuantitativeDataset) -> ReportPlan:
    return ReportPlan(
        title="Report Title",
        subtitle="Report Subtitle",
        sections=[
            ReportSection(heading="Summary", narrative=[ContentItem(text="bullet"), ContentItem(text="prose", bullet=False)]),
            ReportSection(heading="Chart", chart=ChartSpec(dataset_id=ds.dataset_id, chart_type=ChartType.PIE, category_column="segment", value_columns=["revenue"], notes="caption")),
            ReportSection(heading="Table", table=TableSpec(dataset_id=ds.dataset_id)),
            ReportSection(heading="Divider", layout_hint=SectionLayoutHint.SECTION_HEADER),
            ReportSection(heading="Appendix note", appendix=True, narrative=[ContentItem(text="methods")]),
        ],
    )


def test_build_docx_has_title_headings_table_and_image(tmp_path):
    ds = _dataset()
    out = tmp_path / "report.docx"

    build_docx(_plan(ds), [ds], LayoutInstructions(include_appendix=True), out)

    doc = Document(str(out))
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading") or p.style.name == "Title"]
    assert "Report Title" in headings
    assert "Summary" in headings
    assert "Appendix" in headings
    assert headings.index("Appendix") < headings.index("Appendix note")
    assert len(doc.tables) == 1
    assert doc.tables[0].cell(0, 0).text == "segment"
    # the pie chart was embedded as an image (InlineShapes)
    assert len(doc.inline_shapes) == 1


def test_plan_pdf_is_a_nonempty_pdf_file(tmp_path):
    ds = _dataset()
    out = tmp_path / "report.pdf"

    asyncio.run(write_pdf(render_deck_html(deck_from_plan(_plan(ds)), [ds]), out))  # the path ReportingService takes for pdf

    assert out.exists()
    assert out.read_bytes().startswith(b"%PDF")
    assert out.stat().st_size > 500
