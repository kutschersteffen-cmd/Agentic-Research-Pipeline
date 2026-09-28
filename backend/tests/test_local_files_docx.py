import asyncio

import docx

from arp.grounding import ground_citations
from arp.ingestion.local_files import LocalFileDocumentSource, parse_file_to_text_with_pages
from arp.schemas.common import Citation, CompanyRef, DocType


def _write_docx(path):
    document = docx.Document()
    document.add_heading("Climate Transition Strategy", 1)
    document.add_paragraph("We have set a company-wide net zero target for 2050.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Scope", "tCO2e"
    table.cell(1, 0).text, table.cell(1, 1).text = "Scope 1", "1,234"
    document.save(path)


def test_docx_parses_to_markdown_with_heading_text_and_table(tmp_path):
    path = tmp_path / "strategy.docx"
    _write_docx(path)

    text, page_breaks = parse_file_to_text_with_pages(path)

    assert "## Climate Transition Strategy" in text
    assert "company-wide net zero target for 2050" in text
    assert "| Scope 1" in text and "1,234" in text, "table rows keep their cells together"
    assert page_breaks == [], "a .docx has no fixed pages"


def test_docx_in_the_documents_folder_is_fetched_and_citations_ground(tmp_path):
    folder = tmp_path / "docs" / "acme" / DocType.SUSTAINABILITY_REPORT.value
    folder.mkdir(parents=True)
    _write_docx(folder / "strategy.docx")

    docs = asyncio.run(LocalFileDocumentSource(tmp_path / "docs").fetch(CompanyRef(company_id="acme", name="Acme")))

    assert len(docs) == 1
    citation = Citation(doc_id=docs[0].doc_id, doc_type=docs[0].doc_type, quote="company-wide net zero target for 2050")
    [grounded] = ground_citations([citation], {docs[0].doc_id: docs[0]}, 0.92)
    assert grounded.grounded is True and grounded.page is None
