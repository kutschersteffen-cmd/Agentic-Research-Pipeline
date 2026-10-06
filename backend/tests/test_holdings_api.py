from __future__ import annotations

import hashlib
import io
import json
from datetime import date

import httpx
import openpyxl
import pytest
from fastapi.testclient import TestClient
from typer.testing import CliRunner

from arp.api.auth import Principal, current_user
from arp.api.deps import get_portfolio_store, get_run_store, settings_dep
from arp.api.main import app
from arp.config import Settings
from arp.holdings import load
from arp.holdings.api_source import pull_holder
from arp.holdings.intake import IntakeError, holder_status
from arp.schemas.portfolio import HolderConfig
from arp.snapshots.client import SnapshotClient
from arp.snapshots.schema import DatasetEntry, SnapshotManifest
from arp.storage.identifier_map import IdentifierMapStore
from arp.storage.portfolio_store import PortfolioStore
from arp.storage.run_store import RunStore
from tests.conftest import PRINCIPAL

ISIN, ISIN2 = "US0378331005", "DE0007164600"
LEI, BAD_LEI = "5493001KJTIIGC8Y1R12", "5493001KJTIIGC8Y1R13"
ANALYST = Principal(user_id="u_ana", name="Ana", role="analyst")
URL = "/api/holdings"
HEADER = "ISIN,LEI,Name,Weight (%),Market value,Currency\n"


def row(as_of, isin=ISIN, weight=60.0, currency="EUR", fx=1.0, holder="P1"):
    return {"portfolio_id": holder, "issuer_key": LEI, "isin": isin, "as_of": as_of, "issuer_scheme": "LEI",
            "weight": weight, "market_value": weight * 10, "currency": currency, "fx_rate_to_eur": fx,
            "source_file": None}


def month_rows(as_of, **kw):
    return [row(as_of, **kw), row(as_of, isin=ISIN2, weight=40.0)]


def upstream(months: dict[str, list[dict]], fail: set[str] = frozenset(), body: bytes | None = None,
             manifest: dict | None = None) -> SnapshotClient:
    """A fake ARP instance serving one revision per month with matching hashes."""
    def handler(request: httpx.Request) -> httpx.Response:
        parts = request.url.path.split("/")  # "", api, v1, snapshots, month, what
        month, what = parts[4], parts[5]
        if month in fail:
            return httpx.Response(500)
        rows = months[month]
        data = body if body is not None else "".join(json.dumps(r) + "\n" for r in rows).encode()
        if what == "manifest" and manifest is not None:
            return httpx.Response(200, json=manifest)
        if what == "manifest":
            m = SnapshotManifest(
                snapshot_id=f"{month}.r1", month=month, revision=1, as_of=rows[0]["as_of"], frozen_at="t",
                schema_version="1.0", datasets=[DatasetEntry(name="portfolio_holdings", major=1, schema_version="1.0",
                                                             rows=len(rows), files={"jsonl": hashlib.sha256(data).hexdigest()})])
            return httpx.Response(200, json=m.model_dump())
        return httpx.Response(200, content=data)

    return SnapshotClient("https://up.example", None, http=httpx.Client(transport=httpx.MockTransport(handler)))


@pytest.fixture
def env(tmp_path):
    store, rs, idmap = PortfolioStore(tmp_path / "pf"), RunStore(tmp_path / "runs"), IdentifierMapStore(tmp_path / "id.jsonl")
    store.save_holder(HolderConfig(holder_id="P1", kind="portfolio", source="api"))
    yield store, rs, idmap
    for dep in (get_portfolio_store, get_run_store, settings_dep, current_user):
        app.dependency_overrides.pop(dep, None)


def pull(env, client, as_of, today=date(2026, 12, 5)):
    store, rs, idmap = env
    holder = store.get_holder("portfolio", "P1")
    return pull_holder(holder, as_of, client=client, base_url="https://up.example", store=store,
                       idmap=idmap, today=today)


def api(env, tmp_path, who=PRINCIPAL, **settings):
    store, rs, _ = env
    app.dependency_overrides[get_portfolio_store] = lambda: store
    app.dependency_overrides[get_run_store] = lambda: rs
    app.dependency_overrides[settings_dep] = lambda: Settings(identifier_map_path=tmp_path / "id.jsonl", **settings)
    app.dependency_overrides[current_user] = lambda: who
    return TestClient(app)


def upload(c, body: bytes, name="positions.csv", **form):
    data = {"holder_id": "P1", "kind": "portfolio", "as_of": "2026-09-30", **form}
    return c.post(f"{URL}/upload", data=data, files={"file": (name, body)})


def test_fake_api_month_lands_once_when_pulled_twice(env):
    client = upstream({"2026-10": month_rows("2026-10-31")})
    assert pull(env, client, "2026-10-31").status == "written"
    second = pull(env, client, "2026-10-31")
    assert env[0].list_revisions("portfolio", "P1", "2026-10-31") == [1]
    assert second.status == "unchanged"
    holder = env[0].get_holder("portfolio", "P1")
    assert holder.last_pull_at and holder.last_error is None and holder.as_of == "2026-10-31"


def test_failed_pull_keeps_last_month_and_flags(env):
    pull(env, upstream({"2026-10": month_rows("2026-10-31")}), "2026-10-31")
    with pytest.raises(httpx.HTTPStatusError):
        pull(env, upstream({}, fail={"2026-11"}), "2026-11-30")
    store = env[0]
    assert store.get_holder("portfolio", "P1").last_error
    assert {h.isin for h in load(store, "portfolio", "P1", "2026-11-30")} == {ISIN, ISIN2}
    assert holder_status(store, date(2026, 12, 5))[0]["stale"] is True


def test_non_eur_api_month_lands_with_fx_rate(env):
    pull(env, upstream({"2026-10": month_rows("2026-10-31", currency="USD", fx=0.9)}), "2026-10-31")
    h = next(h for h in env[0].load_snapshot("P1", "2026-10-31", kind="portfolio") if h.isin == ISIN)
    assert (h.currency, h.fx_rate_to_eur, h.market_value_eur) == ("USD", 0.9, 540.0)


def test_pull_for_file_holder_refused(env):
    env[0].save_holder(HolderConfig(holder_id="P1", kind="portfolio", source="file"))
    with pytest.raises(IntakeError) as exc:
        pull(env, upstream({"2026-10": month_rows("2026-10-31")}), "2026-10-31")
    assert exc.value.status == 409


def test_upload_bad_lei_422_names_row(env, tmp_path):
    body = (HEADER + f"{ISIN},{LEI},Apple,60,600,EUR\n{ISIN2},{BAD_LEI},SAP,40,400,EUR\n").encode()
    r = upload(api(env, tmp_path), body)
    assert r.status_code == 422
    err = r.json()["detail"]["errors"][0]
    assert (err["row"], err["column"]) == (3, "lei") and err["message"]


def test_upload_writes_and_ignores_formulas(env, tmp_path):
    body = (HEADER + f"{ISIN},,=HYPERLINK(\"x\"),60,600,EUR\n{ISIN2},,SAP,40,400,EUR\n").encode()
    r = upload(api(env, tmp_path), body, override_reason="desk correction")
    assert r.status_code == 200, r.text
    assert r.json()["status"] == "written" and r.json()["revision"] == 1


def test_upload_over_size_limit_413(env, tmp_path):
    r = upload(api(env, tmp_path, max_upload_bytes=10), (HEADER * 2).encode())
    assert r.status_code == 413


def test_upload_unreadable_file_422(env, tmp_path):
    c = api(env, tmp_path)
    for r in (upload(c, b"not a zip", name="positions.xlsx"), upload(c, HEADER.encode(), provider="nope"),
              upload(c, HEADER.encode(), as_of="20260930")):
        assert r.status_code == 422
        assert r.json()["detail"]["errors"] == []


def test_upload_over_api_month_without_reason_409(env, tmp_path):
    pull(env, upstream({"2026-09": month_rows("2026-09-30")}), "2026-09-30")
    body = (HEADER + f"{ISIN},,Apple,50,600,EUR\n{ISIN2},,SAP,50,400,EUR\n").encode()
    assert upload(api(env, tmp_path), body).status_code == 409


def test_template_download_xlsx_opens(env, tmp_path):
    r = api(env, tmp_path).get(f"{URL}/template", params={"kind": "portfolio", "format": "xlsx"})
    assert r.status_code == 200
    assert r.headers["content-disposition"] == 'attachment; filename="holdings-template-portfolio.xlsx"'
    header = next(openpyxl.load_workbook(io.BytesIO(r.content)).active.iter_rows(values_only=True))
    assert header[0] == "ISIN" and "FX rate to EUR" in header


def test_put_holder_needs_approver_and_is_audited(env, tmp_path):
    body = {"name": "Fund", "source": "file"}
    assert api(env, tmp_path, who=ANALYST).put(f"{URL}/holders/portfolio/P1", json=body).status_code == 403
    r = api(env, tmp_path).put(f"{URL}/holders/portfolio/P1", json=body)
    assert r.status_code == 200 and r.json()["source"] == "file" and "user_id" not in r.text
    audit = env[0]._read_jsonl(env[0].holdings_audit_path())[-1]
    assert (audit["user_id"], audit["old_source"], audit["new_source"]) == (PRINCIPAL.user_id, "api", "file")
    assert api(env, tmp_path).put(f"{URL}/holders/portfolio/P1", json={"source": "file"}).json()["name"] == "Fund"


def test_holders_response_has_no_user_id(env, tmp_path):
    c = api(env, tmp_path)
    body = (HEADER + f"{ISIN},,Apple,60,600,EUR\n{ISIN2},,SAP,40,400,EUR\n").encode()
    assert upload(c, body, override_reason="r").status_code == 200
    r = c.get(f"{URL}/holders")
    assert r.json()["holders"][0]["holder_id"] == "P1"
    assert "user_id" not in r.text and PRINCIPAL.user_id not in r.text


def test_pull_endpoint_503_without_api_url(env, tmp_path):
    r = api(env, tmp_path, holdings_api_url=None).post(f"{URL}/holders/portfolio/P1/pull")
    assert r.status_code == 503 and "holdings_api_url is not set" in r.text


def test_cli_import_xlsx(env, tmp_path, monkeypatch):
    from arp.cli import holdings as cli

    users = tmp_path / "users.json"
    users.write_text(json.dumps({"users": [{"user_id": "u_apr", "name": "A", "role": "approver", "token": "ta"}]}))
    monkeypatch.setattr(cli, "get_settings", lambda: Settings(
        users_file=users, portfolios_dir=tmp_path / "pf", runs_dir=tmp_path / "runs",
        identifier_map_path=tmp_path / "id.jsonl", portfolio_backend="file"))
    monkeypatch.setenv("ARP_CLI_TOKEN", "ta")
    wb = openpyxl.Workbook()
    for line in (HEADER.strip().split(","), [ISIN, LEI, "Apple", 60, 600, "EUR"], [ISIN2, None, "SAP", 40, 400, "EUR"]):
        wb.active.append(line)
    path = tmp_path / "positions.xlsx"
    wb.save(path)
    env[0].save_holder(HolderConfig(holder_id="P1", kind="portfolio", source="file"))
    res = CliRunner().invoke(cli.holdings_app, ["import", "--file", str(path), "--holder", "P1", "--as-of", "2026-09-30"])
    assert res.exit_code == 0, res.output
    assert env[0].list_revisions("portfolio", "P1", "2026-09-30") == [1]
    bad = CliRunner().invoke(cli.holdings_app, ["import", "--file", str(path), "--holder", "P1", "--as-of", "2026-13-01"])
    assert bad.exit_code == 1


def _pull_api(env, tmp_path, monkeypatch, client):
    from arp.api.routers import holdings as router

    monkeypatch.setattr(router, "SnapshotClient", lambda *a, **k: client)
    return api(env, tmp_path, holdings_api_url="https://up.example")


class NoNetwork(SnapshotClient):
    def __init__(self):
        super().__init__("https://up.example", None)

    def _get(self, *a, **k):
        raise AssertionError("no network call expected")


def test_pull_endpoint_refuses_bad_as_of_and_file_holder_before_network(env, tmp_path, monkeypatch):
    c = _pull_api(env, tmp_path, monkeypatch, NoNetwork())
    assert c.post(f"{URL}/holders/portfolio/P1/pull", params={"as_of": "20261031"}).status_code == 422
    env[0].save_holder(HolderConfig(holder_id="P1", kind="portfolio", source="file"))
    assert c.post(f"{URL}/holders/portfolio/P1/pull", params={"as_of": "2026-09-30"}).status_code == 409
    holder = env[0].get_holder("portfolio", "P1")
    assert holder.last_error is None and holder.last_pull_at is None


def test_pull_endpoint_malformed_upstream_502(env, tmp_path, monkeypatch):
    rows = {"2026-09": month_rows("2026-09-30")}
    for client in (upstream(rows, body=b"[1, 2]\n"), upstream(rows, body=b"{not json\n"),
                   upstream(rows, manifest={"snapshot_id": 5})):
        r = _pull_api(env, tmp_path, monkeypatch, client).post(f"{URL}/holders/portfolio/P1/pull",
                                                               params={"as_of": "2026-09-30"})
        assert r.status_code == 502 and "upstream" in r.text, r.text
