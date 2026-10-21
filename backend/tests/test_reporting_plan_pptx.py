"""A ReportPlan rendered to pptx the way ReportingService does it: deck_from_plan, then the house renderer."""

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from arp.reporting.deck_compat import deck_from_plan
from arp.reporting.house_pptx import build_house_pptx
from arp.schemas.reporting import (
    ChartSpec,
    ChartType,
    ColumnKind,
    ContentItem,
    DatasetColumn,
    QuantitativeDataset,
    ReportPlan,
    ReportSection,
    SectionLayoutHint,
    TableSpec,
)


def _dataset() -> QuantitativeDataset:
    return QuantitativeDataset(
        name="Revenue",
        columns=[DatasetColumn(name="segment", kind=ColumnKind.CATEGORY), DatasetColumn(name="revenue", kind=ColumnKind.NUMBER)],
        rows=[{"segment": "EV", "revenue": 120}, {"segment": "Grid", "revenue": 80}],
    )


def _plan(ds: QuantitativeDataset) -> ReportPlan:
    return ReportPlan(
        title="Deck Title",
        subtitle="Deck Subtitle",
        sections=[
            ReportSection(heading="Text only", narrative=[ContentItem(text="a"), ContentItem(text="b", bullet=False)]),
            ReportSection(
                heading="Chart section",
                chart=ChartSpec(dataset_id=ds.dataset_id, chart_type=ChartType.COLUMN, category_column="segment", value_columns=["revenue"], notes="EUR m"),
                narrative=[ContentItem(text="context")],
            ),
            ReportSection(heading="Table section", table=TableSpec(dataset_id=ds.dataset_id)),
            ReportSection(heading="Divider", layout_hint=SectionLayoutHint.SECTION_HEADER),
            ReportSection(heading="Appendix note", appendix=True, narrative=[ContentItem(text="methodology")]),
        ],
    )


def _render(tmp_path):
    ds = _dataset()
    out = tmp_path / "deck.pptx"
    build_house_pptx(deck_from_plan(_plan(ds)), [ds], out)
    return Presentation(str(out))


def _texts(slide) -> list[str]:
    return [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame]


def test_plan_deck_has_a_title_slide_then_one_slide_per_section(tmp_path):
    prs = _render(tmp_path)

    # title slide + 5 sections
    assert len(prs.slides._sldIdLst) == 6
    assert "Deck Title" in _texts(prs.slides[0])


def test_plan_deck_chart_section_has_native_chart_shape(tmp_path):
    prs = _render(tmp_path)

    chart_slide = next(s for s in prs.slides if "Chart section" in _texts(s))
    assert any(shape.shape_type == MSO_SHAPE_TYPE.CHART for shape in chart_slide.shapes)


def test_plan_deck_table_section_has_table_shape(tmp_path):
    prs = _render(tmp_path)

    table_slide = next(s for s in prs.slides if "Table section" in _texts(s))
    table_shapes = [shape for shape in table_slide.shapes if shape.has_table]
    assert len(table_shapes) == 1
    table = table_shapes[0].table
    assert table.cell(0, 0).text.lower() == "segment"
    assert table.cell(1, 0).text == "EV"
