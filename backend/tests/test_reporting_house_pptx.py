import shutil
import subprocess

import pytest
from pptx import Presentation
from pypdf import PdfReader

from arp.reporting.browser import write_pdf
from arp.reporting.house_pptx import build_house_pptx
from arp.reporting.house_style import get_variant, load_tokens, slot_rect
from arp.reporting.html_render import render_deck_html
from arp.schemas.reporting import ChartSpec, ColumnKind, DatasetColumn, Deck, QuantitativeDataset, SlideContent, TableSpec
from tests.fixtures.stress_deck import stress_deck

_PX = 9525


def _shapes(slide, name):
    return [s for s in slide.shapes if s.name == name]


def _rgb(font) -> str:
    return str(font.color.rgb).lower()


def test_pptx_positions_match_grid(tmp_path):
    deck, ds = stress_deck("min")
    prs = Presentation(build_house_pptx(deck, ds, tmp_path / "d.pptx"))
    assert (prs.slide_width, prs.slide_height) == (1920 * _PX, 1080 * _PX)
    for slide, content in zip(prs.slides, deck.slides, strict=True):
        for spec in get_variant(content.layout, content.variant).slots:
            r = slot_rect(load_tokens(), spec)
            shape = next(s for s in slide.shapes if s.name == f"slot:{spec.name}")
            assert (shape.left, shape.top, shape.width, shape.height) == tuple(round(v * _PX) for v in r)


def test_pptx_charts_are_native(tmp_path):
    deck, ds = stress_deck("min")
    prs = Presentation(build_house_pptx(deck, ds, tmp_path / "d.pptx"))
    i = next(i for i, s in enumerate(deck.slides) if s.layout == "chart_takeaway")
    assert _shapes(prs.slides[i], "slot:chart")[0].has_chart


def test_pptx_non_native_chart_is_a_picture(tmp_path):
    deck, ds = stress_deck("min")
    slide = next(s for s in deck.slides if s.layout == "chart_takeaway")
    slide.chart = ChartSpec(dataset_id="stress", chart_type="waterfall", category_column="cat", value_columns=["a"])
    prs = Presentation(build_house_pptx(Deck(title="T", slides=[slide]), ds, tmp_path / "d.pptx"))
    shape = _shapes(prs.slides[0], "slot:chart")[0]
    assert shape.shape_type == 13 and not getattr(shape, "has_chart", False)  # MSO_SHAPE_TYPE.PICTURE


def test_pptx_escapes_nothing_and_keeps_literal_text(tmp_path):
    deck = Deck(title="T", slides=[SlideContent(headline="<script> & {{x}}", layout="bullets", variant="three", slots={"items": ["a <b> & {{y}}"]},
                                                speaker_notes="say {{this}}", source_refs=["s1", "s2"])])
    prs = Presentation(build_house_pptx(deck, [], tmp_path / "d.pptx"))
    s = prs.slides[0]
    assert _shapes(s, "headline")[0].text_frame.text == "<script> & {{x}}"
    assert _shapes(s, "slot:items")[0].text_frame.text == "a <b> & {{y}}"
    assert s.notes_slide.notes_text_frame.text == "say {{this}}"
    assert _shapes(s, "footer:refs")[0].text_frame.text == "s1 · s2" and _shapes(s, "footer:page")[0].text_frame.text == "1"


@pytest.mark.parametrize("mode", ["light", "dark"])
def test_pptx_follows_mode_and_metric_numbers_are_ink(tmp_path, mode):
    t = load_tokens()
    c = t.colors(mode)
    deck = Deck(title="T", slides=[SlideContent(headline="h", layout="big_number", variant="two",
                                                slots={"number_1": "1", "label_1": "a", "number_2": "2", "label_2": "b"})])
    s = Presentation(build_house_pptx(deck, [], tmp_path / "d.pptx", mode=mode)).slides[0]
    font = lambda n: _shapes(s, n)[0].text_frame.paragraphs[0].runs[0].font  # noqa: E731
    assert _rgb(font("slot:number_1")) == _rgb(font("slot:number_2")) == c.ink.lstrip("#").lower()
    assert font("slot:label_1").name == t.fonts.mono.split(",")[0].strip("'\" ")
    assert str(s.background.fill.fore_color.rgb).lower() == c.background.lstrip("#").lower()
    assert font("headline").name == t.heading_font(mode).split(",")[0].strip("'\" ")
    assert (font("headline").size.pt, font("slot:number_1").size.pt) == (t.type["headline"].size * 0.75, t.type["big_number"].size * 0.75)
    rule = _shapes(s, "accent_rule")[0]
    assert (rule.left, rule.top) == (t.grid.margin_x * _PX, (t.grid.margin_top - 32 - t.grid.rule_height) * _PX)


def _table_deck(n_rows, **spec):
    ds = QuantitativeDataset(dataset_id="d", name="D", columns=[DatasetColumn(name="k", kind=ColumnKind.CATEGORY), DatasetColumn(name="v")],
                             rows=[{"k": f"r{i}", "v": i} for i in range(n_rows)])
    return Deck(title="T", slides=[SlideContent(headline="h", layout="table", variant="compact", table=TableSpec(dataset_id="d", **spec))]), [ds]


def test_pptx_table_window_and_more_rows_caption(tmp_path):
    deck, ds = _table_deck(10, max_rows=3, row_offset=2)
    s = Presentation(build_house_pptx(deck, ds, tmp_path / "d.pptx")).slides[0]
    table = _shapes(s, "slot:table")[0].table
    assert [table.cell(r, 0).text for r in range(4)] == ["K", "r2", "r3", "r4"]
    assert _shapes(s, "slot:table:more")[0].text_frame.text == "+5 more rows"
    deck, ds = _table_deck(5, max_rows=3, row_offset=2)
    assert not _shapes(Presentation(build_house_pptx(deck, ds, tmp_path / "e.pptx")).slides[0], "slot:table:more")


def test_html_table_more_rows_caption():
    deck, ds = _table_deck(10, max_rows=3, row_offset=2)
    assert "+5 more rows" in render_deck_html(deck, ds)
    deck, ds = _table_deck(5, max_rows=3, row_offset=2)
    assert "more rows" not in render_deck_html(deck, ds).split("<body ", 1)[1]


async def test_html_table_with_more_rows_still_fits():
    from arp.reporting.browser import measure

    deck, ds = _table_deck(40, max_rows=8)
    assert [f for f in await measure(render_deck_html(deck, ds)) if f.stage == "fit"] == []


def _first_words(pdf: str, words: list[str]) -> dict[str, tuple[float, float]]:
    """Top-left (px) of the text run containing each word; the PDFs are 1440pt wide for the 1920px canvas."""
    page, found = PdfReader(pdf).pages[0], {}

    def visit(text, cm, tm, _font, _size):
        for w in words:
            if text.strip().startswith(w) and w not in found:
                x = tm[4] * cm[0] + tm[5] * cm[2] + cm[4]
                y = tm[4] * cm[1] + tm[5] * cm[3] + cm[5]
                found[w] = (x / 0.75, (float(page.mediabox.height) - y) / 0.75)

    page.extract_text(visitor_text=visit)
    return found


@pytest.mark.skipif(shutil.which("soffice") is None, reason="needs LibreOffice")
async def test_pptx_matches_pdf_text_positions(tmp_path):
    deck = Deck(title="T", slides=[SlideContent(headline="Alpha headline", layout="two_column", variant="text_text",
                                                slots={"left": "Bravo left column", "right": "Charlie right column"})])
    pptx = build_house_pptx(deck, [], tmp_path / "d.pptx")
    html_pdf = await write_pdf(render_deck_html(deck, []), tmp_path / "html.pdf")
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(tmp_path), str(pptx)], check=True, capture_output=True, timeout=180)
    words = ["Alpha", "Bravo", "Charlie"]
    a, b = _first_words(str(html_pdf), words), _first_words(str(tmp_path / "d.pdf"), words)
    assert set(a) == set(b) == set(words)
    for w in words:
        assert abs(a[w][0] - b[w][0]) <= 12 and abs(a[w][1] - b[w][1]) <= 12, (w, a[w], b[w])


def test_pptx_native_chart_text_is_projector_sized(tmp_path):
    deck, ds = stress_deck("min")
    slide = next(s for s in deck.slides if s.layout == "chart_takeaway")
    slide.chart = slide.chart.model_copy(update={"title": "A title"})
    prs = Presentation(build_house_pptx(Deck(title="T", slides=[slide]), ds, tmp_path / "d.pptx"))
    chart = _shapes(prs.slides[0], "slot:chart")[0].chart
    sizes = [int(v) for v in chart._chartSpace.xpath(".//@sz")]
    assert sizes and min(sizes) >= 1800  # 18pt = 24px on the 1920px canvas


def test_pptx_image_is_letterboxed_inside_its_slot(tmp_path):
    from PIL import Image

    img = tmp_path / "wide.png"
    Image.new("RGB", (400, 100), "red").save(img)
    deck = Deck(title="T", slides=[SlideContent(headline="h", layout="image", variant="default", slots={"caption": "c"}, image_path=str(img))])
    pic = _shapes(Presentation(build_house_pptx(deck, [], tmp_path / "d.pptx")).slides[0], "slot:image")[0]
    spec = next(s for s in get_variant("image", "default").slots if s.name == "image")
    r = slot_rect(load_tokens(), spec)
    assert (pic.left, pic.top, pic.width, pic.height) == tuple(round(v * _PX) for v in r)
    shown_w = pic.width / (1 - pic.crop_left - pic.crop_right)
    shown_h = pic.height / (1 - pic.crop_top - pic.crop_bottom)
    assert shown_w / shown_h == pytest.approx(4, rel=1e-3)  # aspect kept; the short side is padded, not stretched
    assert pic.crop_top < 0 and pic.crop_left == 0


def test_pptx_title_and_section_show_the_title_once(tmp_path):
    deck = Deck(title="T", slides=[SlideContent(headline="Deck", layout="title", variant="plain", slots={"title": "Deck", "subtitle": "S"}),
                                   SlideContent(headline="Part", layout="section", variant="default", slots={"title": "Part"})])
    prs = Presentation(build_house_pptx(deck, [], tmp_path / "d.pptx"))
    for slide, text in zip(prs.slides, ["Deck", "Part"], strict=True):
        assert not _shapes(slide, "headline") and [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame].count(text) == 1


def test_pptx_table_header_is_small_uppercase_mono(tmp_path):
    deck, ds = _table_deck(3, max_rows=3)
    table = _shapes(Presentation(build_house_pptx(deck, ds, tmp_path / "d.pptx")).slides[0], "slot:table")[0].table
    run = table.cell(0, 0).text_frame.paragraphs[0].runs[0]
    assert run.text == "K" and run.font.name == load_tokens().fonts.mono.split(",")[0].strip("'\" ") and not run.font.bold
    assert run._r.rPr.get("spc")


def test_pptx_renders_every_v2_variant(tmp_path):
    deck, ds = stress_deck("min")
    assert len(Presentation(build_house_pptx(deck, ds, tmp_path / "d.pptx")).slides) == len(deck.slides)
    deck, ds = stress_deck("max", "committee")
    assert len(Presentation(build_house_pptx(deck, ds, tmp_path / "e.pptx")).slides) == len(deck.slides)


def test_pptx_cards_are_separate_shapes(tmp_path):
    deck = Deck(title="T", slides=[SlideContent(headline="h", layout="cards", variant="three", slots={"items": ["One: first body", "Two: second", "Three: third"]})])
    s = Presentation(build_house_pptx(deck, [], tmp_path / "d.pptx")).slides[0]
    cards = [_shapes(s, f"slot:items:{i}") for i in range(3)]
    assert all(len(c) == 1 for c in cards)
    assert [p.text for p in cards[0][0].text_frame.paragraphs][-2:] == ["One", "first body"]


def test_pptx_short_text_uses_short_role(tmp_path):
    deck = Deck(title="T", slides=[SlideContent(headline="h", layout="statement", variant="plain", slots={"statement": "Three short words"})])
    s = Presentation(build_house_pptx(deck, [], tmp_path / "d.pptx")).slides[0]
    spec = get_variant("statement", "plain").slots[0]
    assert _shapes(s, "slot:statement")[0].text_frame.paragraphs[0].runs[0].font.size.pt == load_tokens().type[spec.type_role_short].size * 0.75


def test_pptx_steps_have_a_connector_line(tmp_path):
    deck = Deck(title="T", slides=[SlideContent(headline="h", layout="steps", variant="three", slots={"items": ["A: a", "B: b", "C: c"]})])
    s = Presentation(build_house_pptx(deck, [], tmp_path / "d.pptx")).slides[0]
    assert len(_shapes(s, "slot:items:line")) == 1 and all(_shapes(s, f"slot:items:{i}") for i in range(3))


def test_pptx_committee_patterns_are_native_shapes(tmp_path):
    items = {"tree": ["q :: Ok? :: a :: b", "a :: =Yes :: high", "b :: =No :: low"],
             "flow": ["One :: x :: *y", "Two :: z"], "profile": ["A :: 40%", "B :: 80%"]}
    deck = Deck(title="T", slides=[
        SlideContent(headline="h", eyebrow="Eb", layout="tree", variant="default", slots={"items": items["tree"], "takeaway_bar": "Why: because"}),
        SlideContent(headline="h", layout="flow", variant="default", slots={"items": items["flow"], "takeaway_bar": "x"}),
        SlideContent(headline="h", layout="profile", variant="default", slots={"meters": items["profile"], "total": "1", "left": ["P", "a"], "right": ["Q", "b"], "takeaway_bar": "x"}),
    ])
    tree, flow, prof = Presentation(build_house_pptx(deck, [], tmp_path / "d.pptx")).slides
    assert len([sh for sh in tree.shapes if sh.name.startswith("slot:items:node")]) == 3
    assert len([sh for sh in tree.shapes if sh.name.startswith("slot:items:edge") and not sh.name.endswith("label")]) == 2
    assert _shapes(tree, "eyebrow")[0].text_frame.text == "EB" and _shapes(tree, "slot:takeaway_bar")
    assert _shapes(flow, "slot:items:arrow:0")
    col, box = _shapes(flow, "slot:items:0")[0], _shapes(flow, "slot:items:0:box:0")[0]
    assert box.top - col.top < col.height / 2 and col.height == _shapes(flow, "slot:items:1")[0].height  # boxes under the title, equal panels
    assert len([sh for sh in prof.shapes if sh.name.startswith("slot:meters:") and sh.name.endswith(":fill")]) == 2


def test_pptx_malformed_structure_falls_back_to_a_list(tmp_path):
    deck = Deck(title="T", slides=[SlideContent(headline="h", layout="flow", variant="default", slots={"items": ["no separators", "at all"], "takeaway_bar": "x"})])
    s = Presentation(build_house_pptx(deck, [], tmp_path / "d.pptx")).slides[0]
    assert _shapes(s, "slot:items")[0].text_frame.text == "no separators\nat all"


def test_pptx_heat_cells_are_tinted(tmp_path):
    deck, ds = _table_deck(3, max_rows=3)
    deck.slides[0].variant, deck.slides[0].table.heat = "compact", {"v": [1, 2]}
    table = _shapes(Presentation(build_house_pptx(deck, ds, tmp_path / "d.pptx")).slides[0], "slot:table")[0].table
    fills = [str(table.cell(r, 1).fill.fore_color.rgb) for r in (1, 2, 3)]
    assert len(set(fills)) == 3 and str(table.cell(1, 0).fill.fore_color.rgb).lower() == "ffffff"


def _card_body_pt(density, items, tmp_path):
    deck = Deck(title="T", slides=[SlideContent(headline="h", layout="cards", variant="three", slots={"items": items})])
    s = Presentation(build_house_pptx(deck, [], tmp_path / f"{density}.pptx", density=density)).slides[0]
    return _shapes(s, "slot:items:0")[0].text_frame.paragraphs[-1].runs[0].font.size.pt


def test_pptx_committee_prose_is_body_and_present_short_text_steps_up(tmp_path):
    t, spec = load_tokens(), get_variant("cards", "three").slots[0]
    short = ["Two words", "Two words", "Two words"]
    assert _card_body_pt("committee", short, tmp_path) == t.type["body"].size * 0.75
    assert _card_body_pt("present", short, tmp_path) == t.type[spec.type_role_short].size * 0.75


def test_pptx_card_with_longest_item_fits_its_outline(tmp_path):
    from arp.reporting.house_pptx import _Builder

    spec = get_variant("cards", "three").slots[0]
    words = " ".join(["dolore"] * spec.max_words)
    deck = Deck(title="T", slides=[SlideContent(headline="h", layout="cards", variant="three", slots={"items": [f"Title here: {words}"] * 3})])
    t, r = load_tokens(), slot_rect(load_tokens(), spec)
    card = _shapes(Presentation(build_house_pptx(deck, [], tmp_path / "d.pptx", density="present")).slides[0], "slot:items:0")[0]
    b, w = _Builder(t, "light", "present"), (r.w - 2 * t.grid.gutter) / 3 - 2 * t.grid.gutter
    need = 2 * t.grid.gutter + t.type["subhead"].size + 20 + b._est_h("Title here", t.type["subhead"], w) + 12 + b._est_h(words, b._role(spec, "cards", [words]), w)
    assert need <= card.height / _PX <= r.h


async def _html_boxes(html: str, selector: str) -> list[list[float]]:
    from playwright.async_api import async_playwright

    async with async_playwright() as p:
        browser = await p.chromium.launch()
        page = await browser.new_page(viewport={"width": 1920, "height": 1080})
        await page.set_content(html)
        await page.evaluate("document.fonts.ready")
        boxes = await page.evaluate(f"[...document.querySelectorAll({selector!r})].map(e => {{ const r = e.getBoundingClientRect(); return [r.left, r.top, r.width, r.height]; }})")
        await browser.close()
    return boxes


@pytest.mark.parametrize("density", ["present", "committee"])
async def test_pptx_cards_match_the_content_sized_html_stack(tmp_path, density):
    from arp.reporting.house_pptx import _e

    items = ["Review queue: indicators with no grounded citation", "Analyst decision: approve, edit or reject", "Audit trail: reviewer, time and history"]
    deck = Deck(title="T", slides=[SlideContent(headline="h", layout="cards", variant="three", slots={"items": items})])
    s = Presentation(build_house_pptx(deck, [], tmp_path / "d.pptx", density=density)).slides[0]
    html = await _html_boxes(render_deck_html(deck, [], density=density), "[data-slot=items] li")
    for i, (x, y, w, h) in enumerate(html):
        [card] = _shapes(s, f"slot:items:{i}")
        assert abs(card.left - _e(x)) <= _e(2) and abs(card.top - _e(y)) <= _e(2) and abs(card.width - _e(w)) <= _e(2)
        # _est_h keeps 16% width slack for a fallback face (boxes never clip): at most one extra body line.
        body = load_tokens().type["body"]
        assert _e(h) - _e(2) <= card.height <= _e(h + body.size * body.line_height + 2)


@pytest.mark.parametrize("density", ["present", "committee"])
async def test_pptx_steps_match_the_html_per_density(tmp_path, density):
    from arp.reporting.house_pptx import _e

    deck = Deck(title="T", slides=[SlideContent(headline="h", layout="steps", variant="four", slots={"items": ["A: a", "B: b", "C: c", "D: d"]})])
    s = Presentation(build_house_pptx(deck, [], tmp_path / "d.pptx", density=density)).slides[0]
    nodes = await _html_boxes(render_deck_html(deck, [], density=density), "[data-slot=items] li")  # each node sits at its li's left edge
    tops = [_shapes(s, f"slot:items:{i}:node")[0] for i in range(4)]
    [line] = _shapes(s, "slot:items:line")
    if density == "committee":  # vertical rows sharing the body, one node per row, the rule running down through them
        assert len({n.left for n in tops}) == 1 and line.width == 0
        for (x, y, _, h), n in zip(nodes, tops, strict=True):
            assert abs(n.left - _e(x)) <= _e(2) and abs(n.top + n.height / 2 - _e(y + h / 2)) <= _e(2)
    else:  # one horizontal row of nodes on a horizontal rule
        assert len({n.top for n in tops}) == 1 and line.height == 0
        assert all(abs(n.left - _e(x)) <= _e(2) for (x, _, _, _), n in zip(nodes, tops, strict=True))


async def test_pptx_decision_rows_share_the_body_like_the_html(tmp_path):
    from arp.reporting.house_pptx import _e

    deck = Deck(title="T", slides=[SlideContent(headline="h", layout="decisions", variant="default", slots={"items": ["Agree :: the sample", "Review :: the flags"]})])
    s = Presentation(build_house_pptx(deck, [], tmp_path / "d.pptx")).slides[0]
    html = await _html_boxes(render_deck_html(deck, []), "ol.decisions > li")
    for i, (_, y, _, h) in enumerate(html):
        [row] = _shapes(s, f"slot:items:{i}")
        assert abs(row.top - _e(y)) <= _e(2) and abs(row.height - _e(h)) <= _e(2)


@pytest.mark.parametrize("density", ["present", "committee"])
async def test_pptx_card_rows_match_the_html(tmp_path, density):
    from arp.reporting.house_pptx import _e

    items = ["Review queue: indicators with no grounded citation", "Analyst decision: approve, edit or reject", "Audit trail: reviewer, time and history"]
    deck = Deck(title="T", slides=[SlideContent(headline="h", layout="cards", variant="rows", slots={"items": items})])
    s = Presentation(build_house_pptx(deck, [], tmp_path / "d.pptx", density=density)).slides[0]
    for i, (x, y, w, h) in enumerate(await _html_boxes(render_deck_html(deck, [], density=density), "[data-slot=items] li")):
        [row], [rule] = _shapes(s, f"slot:items:{i}"), _shapes(s, f"slot:items:{i}:rule")
        assert abs(rule.top - _e(y)) <= _e(2) and abs(row.top - _e(y)) <= _e(2) and abs(row.height - _e(h)) <= _e(2)
        assert abs(row.left - _e(x + 112)) <= _e(2) and abs(row.left + row.width - _e(x + w)) <= _e(2)
