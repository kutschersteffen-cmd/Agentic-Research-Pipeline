import pytest
from openpyxl import Workbook

from arp.portfolio.constituent_import import import_constituent_files
from arp.storage.portfolio_store import PortfolioStore

HEADER = [
    None,
    "Name",
    "ISIN",
    "Country",
    "Currency",
    "Exchange",
    "Type of Security",
    "Rating",
    "Primary Listing",
    "Industry Classification",
    "Weighting",
]


def _xlsx(tmp_path, rows, title="2026-10-02", isin="IE00TESTFUND"):
    wb = Workbook()
    ws = wb.active
    ws.title = title
    ws["A2"] = "Disclaimer"
    for col, v in enumerate(HEADER, 1):
        ws.cell(row=4, column=col, value=v)
    for i, r in enumerate(rows, 5):
        for col, v in enumerate([None, *r[:2], r[2], r[3], None, r[4], None, None, r[5], r[6]], 1):
            ws.cell(row=i, column=col, value=v)
    path = tmp_path / f"Constituent_{isin}.xlsx"
    wb.save(path)
    return path


# name, isin, country, ccy, type, industry, weight
ROWS = [
    ("Apple", "US0378331005", "Vereinigte Staaten von Amerika", "USD", "Aktien", "Informationstechnologie", 0.5),
    ("Dead Co", "GB00DEAD0001", "Großbritannien (UK)", "GBP", "Aktien", "Finanzen", "N/A"),
    ("Cash EUR", "CASH_EUR", "-", "EUR", "Cash", "-", -0.01),
    ("Index fut", "___ADI2VD0J5", "-", "USD", "Future", "-", 0.02),
    ("Odd", "XS0000000001", "Atlantis", "EUR", "Gizmo", "Industrieunternehmen", 0.1),
]


def test_import(tmp_path):
    store = PortfolioStore(tmp_path / "store")
    path = _xlsx(tmp_path, ROWS)
    s = import_constituent_files(store, [path], 100.0)
    p = s["portfolios"][0]
    assert (p["rows_read"], p["holdings_written"], p["as_of_date"]) == (5, 4, "2026-10-02")
    assert p["skipped_no_weight"] == [{"name": "Dead Co", "isin": "GB00DEAD0001", "weighting": "N/A"}]
    assert p["weight_sum"] == pytest.approx(0.61)
    assert s["untranslated"] == ["Atlantis"] and s["other_asset_classes"] == ["Gizmo"]
    assert store.get_portfolio("dws_ie00testfund").name == "DWS ETF IE00TESTFUND"
    h = {x.security_id: x for x in store.load_snapshot("dws_ie00testfund", "2026-10-02")}
    assert h["US0378331005"].market_value_eur == 50 and h["US0378331005"].weight_pct == 50
    assert h["CASH_EUR"].market_value_eur == -1  # negative kept
    assert store.get_security("CASH_EUR").asset_class == "cash"
    assert store.get_security("___ADI2VD0J5").asset_class == "derivative"
    assert store.get_security("XS0000000001").asset_class == "other"
    c = store.get_company("isin:US0378331005")
    assert (c.country, c.sector) == ("United States", "Information Technology")
    assert store.get_company("isin:CASH_EUR").sector is None
    assert store.get_company("isin:XS0000000001").country == "Atlantis"


def test_idempotent(tmp_path):
    store = PortfolioStore(tmp_path / "store")
    path = _xlsx(tmp_path, ROWS)
    import_constituent_files(store, [path], 100.0)
    first = store.load_snapshot("dws_ie00testfund", "2026-10-02")
    import_constituent_files(store, [path], 100.0)
    assert store.load_snapshot("dws_ie00testfund", "2026-10-02") == first
    assert len(store.list_portfolios()) == 1


@pytest.mark.parametrize("title", ["Sheet1", "02.10.2026"])
def test_bad_sheet_title(tmp_path, title):
    with pytest.raises(ValueError, match="ISO date"):
        import_constituent_files(PortfolioStore(tmp_path / "s"), [_xlsx(tmp_path, ROWS, title=title)], 1.0)


def test_bad_file_name(tmp_path):
    path = _xlsx(tmp_path, ROWS)
    bad = path.rename(tmp_path / "holdings.xlsx")
    with pytest.raises(ValueError, match="Constituent_"):
        import_constituent_files(PortfolioStore(tmp_path / "s"), [bad], 1.0)


def test_notional_tag_is_plain_digits(tmp_path):
    store = PortfolioStore(tmp_path / "store")
    import_constituent_files(store, [_xlsx(tmp_path, ROWS[:1])], 1e8)
    assert "sizing:assumed-notional-eur-100000000" in store.get_portfolio("dws_ie00testfund").tags


def test_missing_security_type_is_not_reported_as_none(tmp_path):
    rows = [("NoType", "XS0000000002", "Japan", "JPY", None, "Energie", 0.1)]
    s = import_constituent_files(PortfolioStore(tmp_path / "store"), [_xlsx(tmp_path, rows)], 1.0)
    assert s["other_asset_classes"] == ["(blank)"]
