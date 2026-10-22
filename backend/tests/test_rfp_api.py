"""rfp-answers-design: library, detection, auto-fill, review and export (the unit rows)."""

from __future__ import annotations

import asyncio
import io
from datetime import date

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select
from typer.testing import CliRunner

from arp.api import run_scheduling
from arp.api.auth import Principal, authorize, current_user
from arp.api.main import app
from arp.cli import app as cli
from arp.cli import rfp as rfp_cli
from arp.db import models as m
from arp.db.session import transaction
from arp.llm.base import LLMUsage
from arp.rfp import store
from arp.rfp.detect import DetectedQuestions, ReportedQuestion
from arp.rfp.draft import Cited, DraftOut
from arp.rfp.library import QuotedPair, QuotedPairs
from tests import rfp_fixtures as fx

TODAY = date(2026, 10, 21)
USERS = {n: Principal(user_id=n, name=n, role=r) for n, r in
         [("alice", "analyst"), ("bob", "approver"), ("carol", "approver"), ("vic", "viewer")]}
A2, A3, A4 = "Please describe your firm.", "What are your assets under management?", "Do you use derivatives for leverage?"
QUESTIONS = [("RFP!A2", A2, "RFP!B2"), ("RFP!A3", A3, "RFP!B3"), ("RFP!A4", A4, "RFP!B4")]
FIRM, AUM = fx.LIB_ANSWERS[0][1], fx.LIB_ANSWERS[1][1]


class Standin:
    """The model, replaced by fixed structured output. `drafts` maps an RFP question to its draft."""

    def __init__(self, pairs=fx.LIB_ANSWERS, questions=QUESTIONS, drafts=None):
        self.pairs, self.questions, self.drafts = pairs, questions, drafts or {}
        self.shown: list[str] = []  # every drafting prompt

    async def complete_structured(self, *, system, prompt, output_model, **_):
        if output_model is QuotedPairs:
            out = QuotedPairs(pairs=[QuotedPair(question=q, answer=a) for q, a in self.pairs])
        elif output_model is DetectedQuestions:
            out = DetectedQuestions(questions=[ReportedQuestion(location=loc, question=q, answer_location=a)
                                               for loc, q, a in self.questions])
        else:
            self.shown.append(prompt)
            out = self.drafts.get(prompt.rsplit("RFP question: ", 1)[1], DraftOut())
        return out, LLMUsage()


@pytest.fixture
def env(pg, monkeypatch):
    """`env.as_(name)`: a client signed in as that user; `env.llm`: the stand-in the next run uses."""
    jobs = []

    class Launcher:
        def launch(self, run_id, job):
            jobs.append(job)

    class Env:
        llm = Standin()

        def as_(self, name):
            who["user"] = USERS[name]
            return client

        def run_jobs(self):
            while jobs:
                asyncio.run(jobs.pop(0)())

    e, who = Env(), {"user": USERS["bob"]}
    app.dependency_overrides.pop(authorize, None)  # the real role check, on the user below
    app.dependency_overrides[current_user] = lambda: who["user"]
    monkeypatch.setattr(store, "today", lambda: TODAY)
    monkeypatch.setattr(run_scheduling, "get_job_launcher", Launcher)
    monkeypatch.setattr(run_scheduling, "get_llm_client", lambda: e.llm)
    monkeypatch.setattr(run_scheduling, "get_verifier_llm_client", lambda: e.llm)
    monkeypatch.setattr(rfp_cli, "_llm", lambda: e.llm)
    client = TestClient(app)
    return e


def _import(tmp_path, data=None, name="lib_past.xlsx"):
    path = tmp_path / name
    path.write_bytes(data or fx.lib_past_xlsx())
    return CliRunner().invoke(cli, ["rfp", "library", "import", str(path)])


def _entries():
    with transaction() as s:
        rows = list(s.scalars(select(m.RfpLibraryEntry).order_by(m.RfpLibraryEntry.id)))
        s.expunge_all()  # readable after the commit
        return rows


def _entry_id(answer):
    return next(e.id for e in _entries() if e.answer == answer)


def _project(env, drafts=None, questions=QUESTIONS, data=None):
    """A project for rfp_in.xlsx, detected and auto-filled by the stand-in; returns its JSON."""
    env.llm = Standin(questions=questions, drafts=drafts)
    r = env.as_("alice").post("/api/rfp/projects", data={"investor": "Pension Fund", "due_date": "2026-11-30"},
                              files={"file": ("rfp_in.xlsx", data or fx.rfp_in_xlsx())})
    assert r.status_code == 200, r.text
    env.run_jobs()
    return env.as_("bob").get(f"/api/rfp/projects/{r.json()['id']}").json()


def _q(project, location):
    return next(q for q in project["questions"] if q["location"] == location)


def _url(project, q, action=""):
    return f"/api/rfp/projects/{project['id']}/questions/{q['id']}{action}"


def _draft(answer, entry_answer, quote):
    return DraftOut(answer=answer, citations=[Cited(entry_id=_entry_id(entry_answer), quote=quote)])


# --- library -------------------------------------------------------------------------------------


def test_library_import(env, tmp_path):  # AC-1, AC-5
    r = _import(tmp_path)
    assert r.exit_code == 0 and r.stdout == "imported 3 entries, dropped 0\n"
    entries = _entries()
    assert [e.answer for e in entries] == [a for _, a in fx.LIB_ANSWERS]
    assert {e.status for e in entries} == {"active"} and {e.review_by for e in entries} == {date(2026, 6, 30)}
    assert entries[1].source_file == "lib_past.xlsx" and entries[1].source_location == "Responses!B3"


def test_library_import_drops_ungrounded_pairs(env, tmp_path):  # AC-2
    env.llm = Standin(pairs=[*fx.LIB_ANSWERS, ("What are your fees?", "0.5% p.a. on all mandates.")])
    r = _import(tmp_path)
    assert r.exit_code == 0 and r.stdout.splitlines() == ["imported 3 entries, dropped 1", "dropped: What are your fees?"]
    assert len(_entries()) == 3


def test_library_review_by_without_document_date(env, tmp_path):  # AC-6
    assert _import(tmp_path, fx.without_core_dates(fx.lib_past_xlsx())).exit_code == 0
    assert {e.review_by for e in _entries()} == {date(2027, 10, 21)}


def test_retire_needs_approver(env, tmp_path):  # AC-7, AC-8
    _import(tmp_path)
    eid = _entries()[0].id
    assert env.as_("alice").post(f"/api/rfp/library/{eid}/retire").status_code == 403
    assert _entries()[0].status == "active"
    r = env.as_("bob").post(f"/api/rfp/library/{eid}/retire")
    assert r.status_code == 200 and r.json()["status"] == "retired" and _entries()[0].status == "retired"


def test_retired_entries_are_never_shown(env, tmp_path):  # AC-9
    env.llm = Standin(pairs=fx.LIB_ANSWERS[:1])
    _import(tmp_path)
    env.as_("bob").post(f"/api/rfp/library/{_entries()[0].id}/retire")
    project = _project(env, drafts={A2: _draft(FIRM, FIRM, FIRM)})
    assert {(q["status"], q["answer"]) for q in project["questions"]} == {("no source", "")}
    assert env.llm.shown == []


# --- detection -----------------------------------------------------------------------------------


def test_detection(env):  # AC-10
    project = _project(env)
    assert [(q["location"], q["answer_location"]) for q in project["questions"]] == [
        ("RFP!A2", "RFP!B2"), ("RFP!A3", "RFP!B3"), ("RFP!A4", "RFP!B4")]
    assert project["not_located"] == 0 and project["run_status"] == "done"


def test_detection_normalises_question(env):  # AC-11
    project = _project(env, questions=[QUESTIONS[0], ("RFP!A3", "what are your assets under management", "RFP!B3"), QUESTIONS[2]])
    assert _q(project, "RFP!A3")["question"] == A3 and project["not_located"] == 0


@pytest.mark.parametrize("questions", [
    [*QUESTIONS, ("RFP!A5", "Any other comments?", "RFP!B5")],  # AC-12: A5 is empty
    [("RFP!A2", A2, "RFP!A6"), *QUESTIONS[1:]],  # AC-13: A6 holds investor text
    [("RFP!A2", A2, "RFP!A2"), *QUESTIONS[1:]],  # the answer cell is the question cell
    [("RFP!A2", A2, "Nope!B2"), *QUESTIONS[1:]],
])
def test_detection_not_located(env, questions):
    project = _project(env, questions=questions)
    assert len(project["questions"]) == len(questions) - 1 and project["not_located"] == 1


def test_detection_word_paragraph_answer_after_itself(env):
    from arp.rfp.detect import detect_questions
    from arp.rfp.locations import Layout

    llm = Standin(questions=[("P1", "Describe your firm.", "P1"), ("P0", "Section A", "P0")])
    found, not_located = asyncio.run(detect_questions(Layout(fx.rfp_para_docx(), ".docx"), llm))
    assert [(d.location, d.answer_location) for d in found] == [("P0", "P0"), ("P1", "P1")] and not_located == 0


# --- auto-fill -----------------------------------------------------------------------------------


def test_autofill_draft_with_sources_and_flags(env, tmp_path):  # AC-14, AC-15
    _import(tmp_path)
    quote = "an independent manager founded in 2004"
    a2 = _q(_project(env, drafts={A2: _draft(FIRM, FIRM, quote)}), "RFP!A2")
    assert a2["status"] == "drafted" and a2["answer"] == FIRM
    assert [(s["entry_id"], s["quote"], s["file"], s["location"]) for s in a2["sources"]] == [
        (_entry_id(FIRM), quote, "lib_past.xlsx", "Responses!B2")]
    assert a2["flags"] == ["check figure", "stale source"]


@pytest.mark.parametrize(("answer", "quote", "status"), [
    ("EUR 13.0bn.", "EUR 13.0bn", "no source"),  # AC-16: the quote is not in the entry
    ("EUR 12.9bn.", "EUR 12.4bn as of 31 Dec 2025", "no source"),  # AC-17: a number no entry has
    ("EUR 12.4bn (31 Dec 2025).", "EUR 12.4bn as of 31 Dec 2025", "drafted"),  # AC-18
])
def test_autofill_checks_quotes_and_numbers(env, tmp_path, answer, quote, status):
    _import(tmp_path)
    a3 = _q(_project(env, drafts={A3: _draft(answer, AUM, quote)}), "RFP!A3")
    assert a3["status"] == status and a3["answer"] == ("" if status == "no source" else answer)


def test_autofill_citation_of_an_unshown_entry_is_no_source(env, tmp_path):
    _import(tmp_path)
    bogus = DraftOut(answer="EUR 12.4bn.", citations=[Cited(entry_id=999, quote="EUR 12.4bn as of 31 Dec 2025")])
    assert _q(_project(env, drafts={A3: bogus}), "RFP!A3")["status"] == "no source"


def test_autofill_never_touches_edited(env, tmp_path):  # AC-19
    _import(tmp_path)
    project = _project(env, drafts={A3: _draft("EUR 12.4bn.", AUM, "EUR 12.4bn as of 31 Dec 2025")})
    a3 = _q(project, "RFP!A3")
    env.as_("alice").put(_url(project, a3), json={"answer": "Our AUM is EUR 12.4bn."})
    env.llm.drafts = {A3: _draft("EUR 12.4bn (31 Dec 2025).", AUM, "EUR 12.4bn as of 31 Dec 2025")}
    assert env.as_("alice").post(f"/api/rfp/projects/{project['id']}/autofill").status_code == 200
    env.run_jobs()
    a3 = _q(env.as_("bob").get(f"/api/rfp/projects/{project['id']}").json(), "RFP!A3")
    assert (a3["answer"], a3["status"]) == ("Our AUM is EUR 12.4bn.", "edited")


# --- review --------------------------------------------------------------------------------------

NO_DERIVATIVES = "No, we do not use derivatives for leverage."


def _a4_written_by(env, user):
    project = _project(env)
    r = env.as_(user).put(_url(project, _q(project, "RFP!A4")), json={"answer": NO_DERIVATIVES})
    return project, _q(r.json(), "RFP!A4")


def test_write_by_hand(env):  # AC-20
    _, a4 = _a4_written_by(env, "alice")
    assert (a4["status"], a4["sources"], a4["flags"], a4["last_editor"]) == ("edited", [], [], "alice")


def test_viewer_cannot_write(env):  # AC-21
    project, a4 = _a4_written_by(env, "alice")
    assert env.as_("vic").put(_url(project, a4), json={"answer": "x"}).status_code == 403
    assert _q(env.as_("vic").get(f"/api/rfp/projects/{project['id']}").json(), "RFP!A4")["answer"] == NO_DERIVATIVES


def test_analyst_cannot_approve(env):  # AC-22
    project, a4 = _a4_written_by(env, "alice")
    assert env.as_("alice").post(_url(project, a4, "/approve"), json={"confirmed_flags": []}).status_code == 403
    assert _q(env.as_("bob").get(f"/api/rfp/projects/{project['id']}").json(), "RFP!A4")["status"] == "edited"


def test_approve_adds_library_entry(env):  # AC-23, AC-24
    project, a4 = _a4_written_by(env, "alice")
    r = env.as_("bob").post(_url(project, a4, "/approve"), json={"confirmed_flags": []})
    assert r.status_code == 200 and _q(r.json(), "RFP!A4")["status"] == "approved"
    [entry] = env.as_("bob").get("/api/rfp/library").json()
    assert (entry["question"], entry["answer"], entry["source"], entry["review_by"], entry["has_figures"], entry["status"]) == (
        A4, NO_DERIVATIVES, "written by alice", "2027-10-21", False, "active")


def test_four_eyes(env):  # AC-47, AC-48
    project, a4 = _a4_written_by(env, "bob")
    assert env.as_("bob").post(_url(project, a4, "/approve"), json={"confirmed_flags": []}).status_code == 403
    assert _q(env.as_("bob").get(f"/api/rfp/projects/{project['id']}").json(), "RFP!A4")["status"] == "edited"
    r = env.as_("carol").post(_url(project, a4, "/approve"), json={"confirmed_flags": []})
    assert r.status_code == 200 and _q(r.json(), "RFP!A4")["status"] == "approved"


def _a2_drafted(env, tmp_path):
    _import(tmp_path)
    project = _project(env, drafts={A2: _draft(FIRM, FIRM, "an independent manager founded in 2004")})
    return project, _q(project, "RFP!A2")


def test_approve_needs_every_flag_confirmed(env, tmp_path):  # AC-25, AC-26
    project, a2 = _a2_drafted(env, tmp_path)
    r = env.as_("bob").post(_url(project, a2, "/approve"), json={"confirmed_flags": ["check figure"]})
    assert r.status_code == 409 and len(_entries()) == 3
    assert _q(env.as_("bob").get(f"/api/rfp/projects/{project['id']}").json(), "RFP!A2")["status"] == "drafted"
    r = env.as_("bob").post(_url(project, a2, "/approve"), json={"confirmed_flags": ["check figure", "stale source"]})
    assert r.status_code == 200 and _q(r.json(), "RFP!A2")["status"] == "approved"
    new = _entries()[3:]
    assert len(new) == 1 and new[0].source_project_id == project["id"] and new[0].has_figures is True


def test_empty_answer_cannot_be_approved(env):  # AC-27
    project = _project(env)
    a4 = _q(project, "RFP!A4")
    assert env.as_("bob").post(_url(project, a4, "/approve"), json={"confirmed_flags": []}).status_code == 409
    assert _q(env.as_("bob").get(f"/api/rfp/projects/{project['id']}").json(), "RFP!A4")["status"] == "no source"
    assert _entries() == []


def test_assign(env):  # AC-28
    project = _project(env)
    for loc in ("RFP!A2", "RFP!A3"):
        assert env.as_("alice").put(_url(project, _q(project, loc), "/assignee"), json={"user_id": "alice"}).status_code == 200
    got = env.as_("alice").get(f"/api/rfp/projects/{project['id']}").json()
    assert [q["assignee"] for q in got["questions"]] == ["alice", "alice", None]


def test_leave_blank(env, tmp_path):  # AC-29
    project, a2 = _a2_drafted(env, tmp_path)
    r = env.as_("alice").post(_url(project, a2, "/leave-blank"))
    assert r.status_code == 200 and _q(r.json(), "RFP!A2")["status"] == "left blank"


def test_edit_keeps_sources_and_recomputes_flags(env, tmp_path):  # AC-50
    project, a2 = _a2_drafted(env, tmp_path)
    r = env.as_("alice").put(_url(project, a2), json={"answer": "Example AM was founded in 2004 and has 85 staff."})
    a2 = _q(r.json(), "RFP!A2")
    assert a2["status"] == "edited" and [s["entry_id"] for s in a2["sources"]] == [_entry_id(FIRM)]
    assert a2["flags"] == ["check figure", "stale source"]


# --- export --------------------------------------------------------------------------------------


def test_export(env, tmp_path):  # AC-38, AC-51
    project, a2 = _a2_drafted(env, tmp_path)
    a3 = _q(project, "RFP!A3")
    env.as_("alice").put(_url(project, a3), json={"answer": "EUR 12.4bn."})
    env.as_("bob").post(_url(project, a2, "/approve"), json={"confirmed_flags": ["check figure", "stale source"]})
    env.as_("bob").post(_url(project, a3, "/approve"), json={"confirmed_flags": ["check figure"]})
    r = env.as_("vic").get(f"/api/rfp/projects/{project['id']}/export")
    assert r.status_code == 200 and r.headers["X-RFP-Unapproved"] == "1"
    before, after = load_workbook(io.BytesIO(fx.rfp_in_xlsx())), load_workbook(io.BytesIO(r.content))
    ws = after["RFP"]
    assert (ws["B2"].value, ws["B3"].value, ws["B4"].value) == (FIRM, "EUR 12.4bn.", None)
    assert after.sheetnames == before.sheetnames
    for row in before["RFP"].iter_rows():
        for c in row:
            if c.coordinate not in ("B2", "B3"):
                new = ws[c.coordinate]
                assert (new.value, new.font.b, new.fill.fgColor.rgb) == (c.value, c.font.b, c.fill.fgColor.rgb)


def test_missing_api_key_leaves_no_project(env, monkeypatch):
    def no_key():
        raise RuntimeError("ARP_ANTHROPIC_API_KEY is not set.")

    monkeypatch.setattr(run_scheduling, "get_llm_client", no_key)
    r = env.as_("alice").post("/api/rfp/projects", data={"investor": "Pension Fund"}, files={"file": ("rfp_in.xlsx", fx.rfp_in_xlsx())})
    assert r.status_code >= 400 and env.as_("alice").get("/api/rfp/projects").json() == []
