import csv
import io
from datetime import date

import openpyxl
import pytest

from arp.holdings.file_source import Mapping, file_ref, load_mapping, read_rows, template
from arp.holdings.validate import RowError, isin_is_valid, validate
from arp.schemas.issuer import lei_is_valid

DEFAULT = load_mapping("default")
HEADER = ["ISIN", "LEI", "Name", "Weight (%)", "Shares", "Free float", "Price", "Currency"]
ISINS = ["US0378331005", "US5949181045", "US88160R1014"]


def _lei(prefix="5493001KJTIIGC8Y1R"):
    for n in range(100):
        cand = f"{prefix}{n:02d}"
        if lei_is_valid(cand):
            return cand
    raise AssertionError


GOOD_LEI = _lei()


def _data(bad_lei=False):
    leis = [GOOD_LEI, "5493001KJTIIGC8Y1R00" if bad_lei else GOOD_LEI, ""]
    assert not (bad_lei and lei_is_valid(leis[1]))
    return [[i, lei, f"N{k}", w, 10, 0.9, 1.5, "EUR"]
            for k, (i, lei, w) in enumerate(zip(ISINS, leis, (50, 30, 20), strict=True))]


def _csv(rows, delim=","):
    buf = io.StringIO()
    csv.writer(buf, delimiter=delim).writerows([HEADER, *rows])
    return buf.getvalue().encode()


def _xlsx(rows):
    wb = openpyxl.Workbook()
    for r in [HEADER, *rows]:
        wb.active.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _v(raw, **kw):
    return validate(raw, kind="index", as_of="2026-09-30", **kw)


def test_csv_and_xlsx_give_identical_rows():
    rows = _data()
    a = _v(read_rows(_csv(rows), "a.csv", DEFAULT))
    b = _v(read_rows(_xlsx(rows), "a.xlsx", DEFAULT))
    assert a.errors == [] and len(a.rows) == 3
    assert a.rows == b.rows


def test_one_bad_lei_rejects_file_and_names_row():
    v = _v(read_rows(_csv(_data(bad_lei=True)), "a.csv", DEFAULT))
    assert v.errors == [RowError(3, "lei", "LEI check digits fail (ISO 17442, mod 97)")]
    assert v.rows == []


def test_isin_check_digit():
    assert isin_is_valid("US0378331005")
    assert not isin_is_valid("US0378331006")


def _w(*ws, **kw):
    raw = [{"_row": i + 2, "isin": ISINS[i], "weight": w} for i, w in enumerate(ws)]
    return _v(raw, **kw).errors


def test_weights_must_sum_to_100():
    assert _w(50, 30, 19) == [RowError(None, "weight", "weights sum to 99%, not 100% ± 0.5")]
    assert _w(50, 30, 19.6) == []
    assert _w(0.5, 0.3, 0.2, weight_unit="fraction") == []


def test_duplicate_position_rejected():
    raw = [{"_row": 2, "isin": ISINS[0], "weight": 50}, {"_row": 3, "isin": ISINS[0], "weight": 50}]
    assert _v(raw).errors == [RowError(3, "isin", "duplicate position")]


def test_future_as_of_rejected():
    raw = [{"_row": 2, "isin": ISINS[0], "weight": 100}]
    v = validate(raw, kind="index", as_of="2026-10-31", today=date(2026, 10, 4))
    assert v.errors == [RowError(None, None, "as-of date is in the future")]


def test_non_eur_portfolio_row_needs_fx_rate():
    raw = [{"_row": 2, "isin": ISINS[0], "weight": 100, "market_value": 5, "currency": "USD"}]
    v = validate(raw, kind="portfolio", as_of="2026-09-30")
    assert v.errors == [RowError(2, "fx_rate_to_eur", "required for a non-EUR position; ARP never invents FX rates")]


def test_decimal_comma_mapping():
    m = Mapping(provider="x", columns=DEFAULT.columns, decimal=",")
    data = b"ISIN;Weight (%)\nUS0378331005;\"100,0\"\n"
    rows = read_rows(data, "a.csv", m)
    v = _v(rows, decimal=m.decimal)
    assert v.errors == [] and v.rows[0]["weight"] == 100.0
    assert parse_check("12,5") == 12.5


def parse_check(s):
    from arp.holdings.validate import parse_decimal
    return parse_decimal(s, ",")


def test_missing_required_column_rejected():
    rows = read_rows(b"ISIN,Name\nUS0378331005,A\n", "a.csv", DEFAULT)
    assert RowError(None, "weight", "missing required column") in _v(rows).errors


def test_template_has_mapped_headers():
    assert template("index", "csv").decode().splitlines()[0] == ",".join(HEADER)
    ws = openpyxl.load_workbook(io.BytesIO(template("index", "xlsx"))).active
    assert [c.value for c in ws[1]] == HEADER


def test_unknown_provider_and_suffix_and_ref():
    with pytest.raises(ValueError):
        load_mapping("../x")
    with pytest.raises(ValueError):
        load_mapping("nope")
    with pytest.raises(ValueError):
        read_rows(b"", "a.txt", DEFAULT)
    assert file_ref(b"a").startswith("sha256:")


@pytest.mark.parametrize("bad", ["nan", "inf", "1e999", float("nan")])
def test_non_finite_numbers_rejected(bad):
    raw = [{"_row": 2, "isin": ISINS[0], "weight": bad}]
    assert RowError(2, "weight", "not a number") in _v(raw).errors
    raw = [{"_row": 2, "isin": ISINS[0], "weight": 100, "market_value": 1, "currency": "USD", "fx_rate_to_eur": bad}]
    v = validate(raw, kind="portfolio", as_of="2026-09-30")
    assert RowError(2, "fx_rate_to_eur", "not a number") in v.errors


def test_unreadable_files_raise_value_error():
    with pytest.raises(ValueError, match="unreadable file"):
        read_rows(b"not a zip", "a.xlsx", DEFAULT)
    wrong = Mapping(provider="x", columns=DEFAULT.columns, sheet="Nope")
    with pytest.raises(ValueError, match="unreadable file"):
        read_rows(_xlsx(_data()), "a.xlsx", wrong)
    with pytest.raises(ValueError, match="too many rows"):
        read_rows(_csv(_data()), "a.csv", DEFAULT, max_rows=2)
    with pytest.raises(ValueError, match="too many rows"):
        read_rows(_xlsx(_data()), "a.xlsx", DEFAULT, max_rows=2)


def test_sign_and_eur_rate_rules():
    def one(**kw):
        raw = [{"_row": 2, "isin": ISINS[0], "weight": 100, "market_value": 1, "currency": "EUR", **kw}]
        return validate(raw, kind="portfolio", as_of="2026-09-30").errors

    assert one(weight=-5) and RowError(2, "weight", "must not be negative") in one(weight=-5)
    assert one(fx_rate_to_eur=0) == [RowError(2, "fx_rate_to_eur", "must be positive")]
    assert one(fx_rate_to_eur=1.1) == [RowError(2, "fx_rate_to_eur", "must be 1 for a EUR position")]
    assert one(fx_rate_to_eur=1) == []


def test_bom_invalid_as_of_and_numeric_name():
    rows = read_rows(b"\xef\xbb\xbf" + _csv(_data()), "a.csv", DEFAULT)
    assert _v(rows).errors == []
    with pytest.raises(ValueError):
        validate(rows, kind="index", as_of="not-a-date")
    r = [{"_row": 2, "isin": ISINS[0], "weight": 100, "name": 123.0}]
    assert _v(r).rows[0]["name"] == "123"
