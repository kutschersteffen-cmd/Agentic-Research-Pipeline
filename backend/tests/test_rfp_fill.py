from __future__ import annotations

import io
import json
from datetime import date

import pytest
from docx import Document
from openpyxl import load_workbook
from typer.testing import CliRunner

from arp.cli import app as cli
from arp.rfp.fill import FillRefused, fill
from arp.rfp.locations import Layout, document_date
from arp.rfp.text import has_figures, is_placeholder, normalise, numbers, plus_12_months
from tests import rfp_fixtures as fx


def test_has_figures():  # AC-3, AC-4
    assert has_figures("We employ 85 people.") is True
    assert has_figures("We have offices in four countries.") is False


def test_text_rules():
    assert numbers("EUR 1,200 and 12.4bn on 31 Dec 2025") == {"1200", "12.4", "31", "2025"}
    assert numbers("1’000 or 1'000") == {"1000"}
    assert is_placeholder("  ") and is_placeholder("[Please answer here]") and not is_placeholder("[a] [b]")
    assert normalise("  What are   your AUM? ") == "what are your aum"
    assert normalise("what are your assets under management") == normalise("What are your assets under management?")
    assert plus_12_months(date(2024, 2, 29)) == date(2025, 2, 28)


def test_layout_order_and_dates():
    xl = Layout(fx.rfp_in_xlsx(), ".xlsx")
    assert [loc for loc, _ in xl.items] == ["RFP!A1", "RFP!B1", "RFP!A2", "RFP!A3", "RFP!A4", "RFP!A6"]
    assert xl.text("RFP!B2") == "" and xl.text("Other!B2") is None and xl.text("RFP!") is None
    word = Layout(fx.rfp_table_docx(), ".docx")
    assert [loc for loc, _ in word.items] == ["P0", "T0R0C0", "T0R0C1", "T0R1C0", "T0R1C1", "P1"]
    assert word.text("T0R9C9") is None and word.text("X1") is None
    assert document_date(fx.lib_past_xlsx()) == date(2025, 6, 30)
    assert document_date(fx.without_core_dates(fx.lib_past_xlsx())) is None


def _run(tmp_path, original: bytes, suffix: str, answers: list[dict]):
    src, ans, out = tmp_path / f"in{suffix}", tmp_path / "answers.json", tmp_path / f"out{suffix}"
    src.write_bytes(original)
    ans.write_text(json.dumps(answers))
    result = CliRunner().invoke(cli, ["rfp", "fill", str(src), str(ans), "--out", str(out)])
    return result, out


def _cells(wb):
    return {
        (ws.title, c.coordinate): (c.value, c.font.b, c.fill.fgColor.rgb)
        for ws in wb.worksheets for row in ws.iter_rows() for c in row
    }


def test_fill_xlsx(pg, tmp_path):  # AC-30
    answers = [{"answer_location": "RFP!B2", "answer": "Example AM is independent."},
               {"answer_location": "RFP!B3", "answer": "EUR 12.4bn."}]
    result, out = _run(tmp_path, fx.rfp_in_xlsx(), ".xlsx", answers)
    assert result.exit_code == 0 and result.stdout == ""
    before, after = load_workbook(io.BytesIO(fx.rfp_in_xlsx())), load_workbook(out)
    assert after.sheetnames == before.sheetnames
    ws = after["RFP"]
    assert ws["B2"].value == "Example AM is independent." and ws["B3"].value == "EUR 12.4bn."
    assert all(ws[r].fill.fgColor.rgb.endswith(fx.YELLOW) for r in ("B2", "B3", "B4"))
    assert ws["B4"].value is None and ws["A1"].font.b
    old, new = _cells(before), _cells(after)
    for key in old.keys() - {("RFP", "B2"), ("RFP", "B3")}:
        assert new[key] == old[key]


def _paras(doc):
    return [(p.text, p.style.name) for p in doc.paragraphs]


def test_fill_table_cell(pg, tmp_path):  # AC-31, AC-32
    for original in (fx.rfp_table_docx(), fx.rfp_table_docx("[Response]")):
        result, out = _run(tmp_path, original, ".docx", [{"answer_location": "T0R1C1", "answer": "EUR 12.4bn."}])
        assert result.exit_code == 0, result.stderr
        before, after = Document(io.BytesIO(original)), Document(out)
        cell = after.tables[0].cell(1, 1)
        assert [(p.text, p.style.name) for p in cell.paragraphs] == [("EUR 12.4bn.", "Normal")]
        assert len(after.paragraphs) == len(before.paragraphs) and len(after.tables) == len(before.tables)
        texts = lambda d: [c.text for row in d.tables[0].rows for c in row.cells]  # noqa: E731
        assert texts(after)[:3] == texts(before)[:3]


def test_fill_after_paragraph(pg, tmp_path):  # AC-33, AC-34
    result, out = _run(tmp_path, fx.rfp_para_docx(), ".docx", [{"answer_location": "P1", "answer": "Example AM is independent."}])
    assert result.exit_code == 0
    assert _paras(Document(out)) == [("Section A", "Heading 1"), ("Describe your firm.", "Normal"),
                                     ("Example AM is independent.", "Body Text"), ("Next section.", "Body Text")]
    result, out = _run(tmp_path, fx.rfp_para_empty_docx(), ".docx", [{"answer_location": "P0", "answer": "Example AM is independent."}])
    assert result.exit_code == 0
    assert _paras(Document(out)) == [("Describe your firm.", "Normal"), ("Example AM is independent.", "Body Text"),
                                     ("Next.", "Normal")]


def test_fill_last_paragraph_and_shift():
    """Several P<n> answers: an insertion after P0 never shifts a later P<n>; after the last paragraph: Normal."""
    out = fill(fx.rfp_para_docx(), ".docx", [("P0", "a0"), ("P2", "a2")])
    assert _paras(Document(io.BytesIO(out))) == [("Section A", "Heading 1"), ("a0", "Normal"),
                                                 ("Describe your firm.", "Normal"), ("Next section.", "Body Text"),
                                                 ("a2", "Normal")]


@pytest.mark.parametrize("make", [fx.rfp_image_xlsx, fx.rfp_chart_xlsx])
def test_fill_refuses_drawings(pg, tmp_path, make):  # AC-35, AC-36
    result, out = _run(tmp_path, make(), ".xlsx", [{"answer_location": "RFP!B2", "answer": "x"}])
    assert result.exit_code == 1 and "RFP" in result.stderr and len(result.stderr.strip().splitlines()) == 1
    assert not out.exists()


@pytest.mark.parametrize(("make", "suffix", "location"), [
    (fx.rfp_in_xlsx, ".xlsx", "RFP!A6"),  # AC-37: investor text
    (fx.rfp_para_docx, ".docx", "P9"),  # AC-49
    (fx.rfp_in_xlsx, ".xlsx", "Nope!B2"),
    (fx.rfp_in_xlsx, ".xlsx", "RFP!"),
    (fx.rfp_table_docx, ".docx", "T0R1C0"),  # holds the question
    (fx.rfp_table_docx, ".docx", "T3R0C0"),
])
def test_fill_refuses_location(pg, tmp_path, make, suffix, location):
    result, out = _run(tmp_path, make(), suffix, [{"answer_location": location, "answer": "x"}])
    assert result.exit_code == 1 and location in result.stderr and len(result.stderr.strip().splitlines()) == 1
    assert not out.exists()


def test_fill_formula_like_answer_stays_text():
    out = fill(fx.rfp_in_xlsx(), ".xlsx", [("RFP!B2", "=1+1")])
    assert load_workbook(io.BytesIO(out))["RFP"]["B2"].value == "=1+1"
    assert load_workbook(io.BytesIO(out))["RFP"]["B2"].data_type == "s"


def test_fill_refusal_is_one_exception():
    with pytest.raises(FillRefused):
        fill(fx.rfp_in_xlsx(), ".xlsx", [("RFP!B2", "ok"), ("RFP!A6", "bad")])
