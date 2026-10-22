"""The spec's RFP fixture files (rfp-answers-design, Acceptance setup notes), built in memory."""

from __future__ import annotations

import io
import re
import zipfile

from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

YELLOW = "FFFF00"
LIB_ANSWERS = [
    ("Describe your firm.", "Example AM is an independent manager founded in 2004."),
    ("What is your total AUM?", "EUR 12.4bn as of 31 Dec 2025."),
    ("Do you sign the UN PRI?", "Yes, we have been a PRI signatory since 2010."),
]


def _save(wb) -> bytes:
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def lib_past_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Responses"
    ws.append(["Question", "Answer"])
    for row in LIB_ANSWERS:
        ws.append(list(row))
    # openpyxl stamps `modified` with the save time, so the property is set in the saved archive
    return _core_xml(_save(wb), lambda xml: re.sub(rb"(<dcterms:modified[^>]*>)[^<]*", rb"\g<1>2025-06-30T00:00:00Z", xml))


def _core_xml(data: bytes, edit) -> bytes:
    src, out = zipfile.ZipFile(io.BytesIO(data)), io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            body = src.read(item.filename)
            dst.writestr(item, edit(body) if item.filename == "docProps/core.xml" else body)
    return out.getvalue()


def without_core_dates(data: bytes) -> bytes:
    """`data` with no `modified` and no `created` core property."""
    return _core_xml(data, lambda xml: re.sub(rb"<dcterms:(created|modified)[^>]*>[^<]*</dcterms:\1>", b"", xml))


def rfp_in_workbook() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "RFP"
    ws["A1"], ws["B1"] = "Question", "Response"
    ws["A1"].font = Font(bold=True)
    ws["A2"] = "Please describe your firm."
    ws["A3"] = "What are your assets under management?"
    ws["A4"] = "Do you use derivatives for leverage?"
    for ref in ("B2", "B3", "B4"):
        ws[ref].fill = PatternFill("solid", fgColor=YELLOW)
    ws["A6"] = "Contact: rfp@investor.example"
    return wb


def rfp_in_xlsx() -> bytes:
    return _save(rfp_in_workbook())


def rfp_image_xlsx() -> bytes:
    from openpyxl.drawing.image import Image
    from PIL import Image as PILImage

    png = io.BytesIO()
    PILImage.new("RGB", (4, 4), "red").save(png, format="PNG")
    wb = rfp_in_workbook()
    wb["RFP"].add_image(Image(io.BytesIO(png.getvalue())), "D2")
    return _save(wb)


def rfp_chart_xlsx() -> bytes:
    from openpyxl.chart import BarChart, Reference

    wb = rfp_in_workbook()
    ws = wb["RFP"]
    ws["E1"], ws["E2"] = 1, 2
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=5, min_row=1, max_row=2))
    ws.add_chart(chart, "G2")
    return _save(wb)


def _docx(build) -> bytes:
    doc = Document()
    build(doc)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def rfp_table_docx(answer_cell_text: str = "") -> bytes:
    def build(doc):
        doc.add_paragraph("Section 1")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = "Question", "Response"
        table.cell(1, 0).text = "What is your total AUM?"
        table.cell(1, 1).paragraphs[0].text = answer_cell_text
        table.cell(1, 1).paragraphs[0].style = doc.styles["Normal"]
        doc.add_paragraph("End")

    return _docx(build)


def rfp_para_docx() -> bytes:
    def build(doc):
        doc.add_paragraph("Section A", style="Heading 1")
        doc.add_paragraph("Describe your firm.", style="Normal")
        doc.add_paragraph("Next section.", style="Body Text")

    return _docx(build)


def rfp_para_empty_docx() -> bytes:
    def build(doc):
        doc.add_paragraph("Describe your firm.", style="Normal")
        doc.add_paragraph("", style="Body Text")
        doc.add_paragraph("Next.", style="Normal")

    return _docx(build)
