import pytest

from arp.reporting.service import ReportingService
from arp.schemas.reporting import (
    ColumnKind,
    ContentItem,
    DatasetColumn,
    LayoutInstructions,
    OutputFormat,
    QuantitativeDataset,
    ReportPlan,
    ReportRequest,
    ReportSection,
    ReportStatus,
)
from arp.storage.reporting_store import ReportingStore


def _store(tmp_path) -> ReportingStore:
    return ReportingStore(tmp_path / "reports", tmp_path / "templates")


def _plan() -> ReportPlan:
    return ReportPlan(title="T", sections=[ReportSection(heading="Summary", narrative=[ContentItem(text="a point")])])


async def test_create_and_plan_persists_manifest_and_plan(tmp_path, fake_llm):
    store = _store(tmp_path)
    service = ReportingService(store)
    llm = fake_llm({"ReportPlan": [_plan()]})
    request = ReportRequest(title="T", qualitative_notes="notes")

    manifest = await service.create_and_plan(request, llm)

    assert manifest.status == ReportStatus.PLAN_READY
    assert store.load_plan(manifest.report_id).title == "T"
    assert store.load_request(manifest.report_id) == request


async def test_create_and_plan_marks_failed_on_llm_error(tmp_path, fake_llm):
    store = _store(tmp_path)
    service = ReportingService(store)
    llm = fake_llm({})  # no scripted response -> FakeLLMClient raises AssertionError
    request = ReportRequest(title="T", qualitative_notes="notes")

    with pytest.raises(AssertionError):
        await service.create_and_plan(request, llm)

    manifest = store.load_manifest(next(m.report_id for m in store.list_reports()))
    assert manifest.status == ReportStatus.FAILED
    assert manifest.error


async def test_run_end_to_end_renders_pptx(tmp_path, fake_llm):
    store = _store(tmp_path)
    service = ReportingService(store)
    llm = fake_llm({"ReportPlan": [_plan()]})
    request = ReportRequest(title="T", qualitative_notes="notes", layout=LayoutInstructions(output_format=OutputFormat.PPTX))

    manifest = await service.run(request, llm)

    assert manifest.status == ReportStatus.COMPLETED
    assert manifest.output_filename == "output.pptx"
    assert store.output_path(manifest.report_id, manifest.output_filename).exists()


async def test_run_end_to_end_renders_docx_with_dataset_chart(tmp_path, fake_llm):
    store = _store(tmp_path)
    service = ReportingService(store)
    ds = QuantitativeDataset(
        name="Revenue", columns=[DatasetColumn(name="segment", kind=ColumnKind.CATEGORY), DatasetColumn(name="revenue", kind=ColumnKind.NUMBER)],
        rows=[{"segment": "EV", "revenue": 100}],
    )
    from arp.schemas.reporting import ChartSpec, ChartType

    plan = ReportPlan(
        title="T",
        sections=[ReportSection(heading="Chart", chart=ChartSpec(dataset_id=ds.dataset_id, chart_type=ChartType.BAR, category_column="segment", value_columns=["revenue"]))],
    )
    llm = fake_llm({"ReportPlan": [plan]})
    request = ReportRequest(title="T", qualitative_notes="notes", datasets=[ds], layout=LayoutInstructions(output_format=OutputFormat.DOCX))

    manifest = await service.run(request, llm)

    assert manifest.status == ReportStatus.COMPLETED
    assert manifest.output_filename == "output.docx"


def test_render_from_plan_raises_when_no_plan_exists(tmp_path):
    store = _store(tmp_path)
    service = ReportingService(store)
    request = ReportRequest(title="T", qualitative_notes="notes")
    from arp.schemas.reporting import ReportManifest

    manifest = ReportManifest(title="T")
    store.save_request(manifest.report_id, request)
    store.save_manifest(manifest)

    with pytest.raises(ValueError, match="No plan drafted"):
        service.render_from_plan(manifest.report_id)


def test_render_from_plan_renders_human_edited_plan(tmp_path):
    """A plan overwritten via ReportingStore.save_plan (simulating a human
    edit through PUT /api/reports/{id}/plan) is what gets rendered, not
    whatever the model originally drafted."""
    store = _store(tmp_path)
    service = ReportingService(store)
    request = ReportRequest(title="T", qualitative_notes="notes", layout=LayoutInstructions(output_format=OutputFormat.PPTX))
    from arp.schemas.reporting import ReportManifest

    manifest = ReportManifest(title="T", output_format=OutputFormat.PPTX)
    store.save_request(manifest.report_id, request)
    store.save_manifest(manifest)
    edited_plan = ReportPlan(title="Edited Title", sections=[ReportSection(heading="Edited Section")])
    store.save_plan(manifest.report_id, edited_plan)

    result = service.render_from_plan(manifest.report_id)

    assert result.status == ReportStatus.COMPLETED
    from pptx import Presentation

    prs = Presentation(str(store.output_path(manifest.report_id, "output.pptx")))
    titles = [s.shapes.title.text for s in prs.slides if s.shapes.title]
    assert "Edited Title" in titles
    assert "Edited Section" in titles
