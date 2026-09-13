from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from arp.reporting.deck_builder import build_deck
from arp.schemas.reporting import (
    ChartSpec,
    ChartType,
    ColumnKind,
    ContentItem,
    DatasetColumn,
    LayoutInstructions,
    QuantitativeDataset,
    ReportPlan,
    ReportSection,
    SectionLayoutHint,
    TableSpec,
)


def _dataset() -> QuantitativeDataset:
    return QuantitativeDataset(
        name="Revenue",
        columns=[DatasetColumn(name="segment", kind=ColumnKind.CATEGORY), DatasetColumn(name="revenue", kind=ColumnKind.NUMBER)],
        rows=[{"segment": "EV", "revenue": 120}, {"segment": "Grid", "revenue": 80}],
    )


def _plan(ds: QuantitativeDataset) -> ReportPlan:
    return ReportPlan(
        title="Deck Title",
        subtitle="Deck Subtitle",
        sections=[
            ReportSection(heading="Text only", narrative=[ContentItem(text="a"), ContentItem(text="b", bullet=False)]),
            ReportSection(
                heading="Chart section",
                chart=ChartSpec(dataset_id=ds.dataset_id, chart_type=ChartType.COLUMN, category_column="segment", value_columns=["revenue"], notes="EUR m"),
                narrative=[ContentItem(text="context")],
            ),
            ReportSection(heading="Table section", table=TableSpec(dataset_id=ds.dataset_id)),
            ReportSection(heading="Divider", layout_hint=SectionLayoutHint.SECTION_HEADER),
            ReportSection(heading="Appendix note", appendix=True, narrative=[ContentItem(text="methodology")]),
        ],
    )


def test_build_deck_produces_one_slide_per_section_plus_title_and_agenda(tmp_path):
    ds = _dataset()
    plan = _plan(ds)
    out = tmp_path / "deck.pptx"

    build_deck(plan, [ds], LayoutInstructions(), None, out)

    prs = Presentation(str(out))
    # title slide + agenda slide + 5 sections
    assert len(prs.slides._sldIdLst) == 7
    assert prs.slides[0].shapes.title.text == "Deck Title"
    assert prs.slides[1].shapes.title.text == "Agenda"


def test_build_deck_chart_section_has_native_chart_shape(tmp_path):
    ds = _dataset()
    plan = _plan(ds)
    out = tmp_path / "deck.pptx"
    build_deck(plan, [ds], LayoutInstructions(), None, out)

    prs = Presentation(str(out))
    chart_slide = next(s for s in prs.slides if s.shapes.title and s.shapes.title.text == "Chart section")
    assert any(shape.shape_type == MSO_SHAPE_TYPE.CHART for shape in chart_slide.shapes)


def test_build_deck_table_section_has_table_shape(tmp_path):
    ds = _dataset()
    plan = _plan(ds)
    out = tmp_path / "deck.pptx"
    build_deck(plan, [ds], LayoutInstructions(), None, out)

    prs = Presentation(str(out))
    table_slide = next(s for s in prs.slides if s.shapes.title and s.shapes.title.text == "Table section")
    table_shapes = [shape for shape in table_slide.shapes if shape.has_table]
    assert len(table_shapes) == 1
    table = table_shapes[0].table
    assert table.cell(0, 0).text == "segment"
    assert table.cell(1, 0).text == "EV"


def test_build_deck_appendix_gets_divider_when_include_appendix(tmp_path):
    ds = _dataset()
    plan = _plan(ds)
    out = tmp_path / "deck.pptx"
    build_deck(plan, [ds], LayoutInstructions(include_appendix=True), None, out)

    prs = Presentation(str(out))
    titles = [s.shapes.title.text for s in prs.slides if s.shapes.title]
    assert "Appendix" in titles
    assert titles.index("Appendix") < titles.index("Appendix note")


def test_build_deck_respects_include_title_and_agenda_false(tmp_path):
    ds = _dataset()
    plan = _plan(ds)
    out = tmp_path / "deck.pptx"
    build_deck(plan, [ds], LayoutInstructions(include_title_slide=False, include_agenda_slide=False), None, out)

    prs = Presentation(str(out))
    titles = [s.shapes.title.text for s in prs.slides if s.shapes.title]
    assert "Deck Title" not in titles
    assert "Agenda" not in titles


def test_build_deck_max_bullets_per_slide_truncates_narrative(tmp_path):
    ds = _dataset()
    plan = ReportPlan(
        title="T",
        sections=[ReportSection(heading="Bullets", narrative=[ContentItem(text=f"point {i}") for i in range(10)])],
    )
    out = tmp_path / "deck.pptx"
    build_deck(plan, [ds], LayoutInstructions(max_bullets_per_slide=3, include_title_slide=False, include_agenda_slide=False), None, out)

    prs = Presentation(str(out))
    slide = prs.slides[0]
    body = next(ph for ph in slide.placeholders if ph.placeholder_format.idx != 0)
    text = body.text_frame.text
    assert "point 0" in text
    assert "point 2" in text
    assert "point 9" not in text
    assert "(+7 more)" in text
